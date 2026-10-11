"""`testler/ornekler/` örnek dosyalarını üretir (Faz 4 / 4g). Dosyalar depoya eklidir; bu betik yalnız yeniden
üretmek içindir:  .venv\\Scripts\\python araclar\\ornek_uret.py

- butce.xlsx   : başka bir programın (openpyxl) yazdığı biçimli çok sayfalı çizelge
- grafikli.xlsx: grafik + gerçek resim + yorum + koşullu biçim + veri doğrulama + tanımlı ad (korunamayan öğeler)
- stok.xlsx    : Faz 5 öğelerinin hepsi (filtre koşullu, doğrulama, koşullu biçim, tanımlı ad, not, ALTTOPLAM)
- grafikler.xlsx: Tabcel'in kendi yazdığı 13 grafik (her tür/alt tür) + resim (Faz 6 / 6j)
- satis_tr.csv : Türkçe Excel'in `;` ayraçlı, Windows-1254 CSV'si
- satis_utf8.csv: virgül ayraçlı UTF-8 (BOM'suz)
"""
import datetime
import os
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.worksheet.filters import CustomFilter, CustomFilters, FilterColumn, Filters
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

KLASOR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "testler", "ornekler")

# 1x1 geçerli PNG (resim parçası taklidi: Pillow gerekmez)
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360f8ffff3f0005fe02fe"
    "a7358a0a0000000049454e44ae426082")


def butce(yol):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bütçe"
    ws.sheet_properties.tabColor = "217346"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Aylık Bütçe 2026"
    ws["A1"].font = Font(name="Calibri", sz=16, b=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="217346")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    basliklar = ["Kalem", "Tür", "Ocak", "Şubat", "Mart", "Toplam"]
    ince = Side(style="thin", color="808080")
    for c, b in enumerate(basliklar, start=1):
        h = ws.cell(row=2, column=c, value=b)
        h.font = Font(b=True)
        h.fill = PatternFill("solid", fgColor="DDEBF7")
        h.border = Border(bottom=Side(style="medium"), top=ince)
        h.alignment = Alignment(horizontal="center")
    veri = [("Maaş", "Gelir", 42000, 42000, 45000), ("Kira", "Gider", 15000, 15000, 15000),
            ("Market", "Gider", 8200.5, 7900.25, 9100), ("Fatura", "Gider", 2300, 2650.75, 2100),
            ("Eğlence", "Gider", 1500, 2200, 1250)]
    for r, satir in enumerate(veri, start=3):
        for c, v in enumerate(satir, start=1):
            h = ws.cell(row=r, column=c, value=v)
            if c >= 3:
                h.number_format = '#,##0.00 "₺"'
        t = ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})")
        t.number_format = '#,##0.00 "₺"'
        t.font = Font(b=True)
    son = 2 + len(veri)
    ws.cell(row=son + 2, column=1, value="Toplam gider").font = Font(i=True)
    ws.cell(row=son + 2, column=6, value='=SUMIF(B3:B7,"Gider",F3:F7)').number_format = '#,##0.00 "₺"'
    ws.cell(row=son + 3, column=1, value="Gider oranı")
    ws.cell(row=son + 3, column=6, value=f"=F{son + 2}/F3").number_format = "0.0%"
    ws.cell(row=son + 4, column=1, value="Son güncelleme")
    ws.cell(row=son + 4, column=6, value=datetime.date(2026, 10, 9)).number_format = "dd.mm.yyyy"
    ws.cell(row=son + 5, column=1, value="Durum")
    ws.cell(row=son + 5, column=6, value=f'=IF(F{son + 3}>0.5,"Dikkat","İyi")')
    ws.cell(row=son + 6, column=1, value="Not")
    ws.cell(row=son + 6, column=2, value="Bu satır uzun bir açıklama metni içerir ve sarılır").alignment = \
        Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 10
    for harf in "CDEF":
        ws.column_dimensions[harf].width = 14
    ws.column_dimensions["H"].hidden = True
    ws.row_dimensions[12].hidden = True
    ws.freeze_panes = "C3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.print_title_rows = "1:2"
    ws.print_options.gridLines = True
    ö = wb.create_sheet("Özet")
    ö["A1"] = "Özet"
    ö["A1"].font = Font(b=True, sz=14)
    ö["A2"] = "Toplam gider"
    ö["B2"] = f"='Bütçe'!F{son + 2}"
    ö["B2"].number_format = '#,##0.00 "₺"'
    ö["A3"] = "Ocak geliri"
    ö["B3"] = "=Bütçe!C3"
    ö["A4"] = "Mantıksal"
    ö["B4"] = True
    ö["A5"] = "Hata"
    ö["B5"] = "=1/0"
    ö.column_dimensions["A"].width = 22
    wb.save(yol)


def grafikli(yol):
    wb = Workbook()
    ws = wb.active
    ws.title = "Satış"
    ws.append(["Ay", "Adet"])
    for i, ay in enumerate(["Ocak", "Şubat", "Mart", "Nisan"], start=1):
        ws.append([ay, i * 10])
    ws["D1"] = "Seçim"
    ws["D2"] = "Evet"
    ws["B2"].comment = Comment("İlk ay değeri", "Ata")
    g = BarChart()
    g.title = "Aylık satış"
    g.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(g, "F2")
    ws.conditional_formatting.add("B2:B5", CellIsRule(operator="greaterThan", formula=["20"],
                                                      fill=PatternFill("solid", bgColor="FFC7CE")))
    dv = DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D2:D10")
    ws.auto_filter.ref = "A1:B5"
    wb.defined_names["SatisToplam"] = DefinedName("SatisToplam", attr_text="Satış!$B$2:$B$5")
    import io
    from openpyxl.drawing.image import Image
    from PIL import Image as PilImage
    png = io.BytesIO()
    PilImage.new("RGB", (120, 60), (68, 114, 196)).save(png, "PNG")            # gerçek resim (6h: artık okunur/yazılır)
    ws.add_image(Image(png), "F20")
    wb.save(yol)


def stok(yol):
    """Stok listesi: Faz 5'in tüm öğeleri (başka bir programın — openpyxl — yazdığı biçimde)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Stok"
    ws.append(["Ürün", "Kategori", "Adet", "Birim fiyat", "Tutar"])
    veri = [("Kulaklık", "Elektronik", 4, 350), ("Şarj aleti", "Elektronik", 25, 120), ("Çay", "Gıda", 60, 45.5),
            ("Kahve", "Gıda", 8, 210), ("Tişört", "Giyim", 30, 150), ("Çorap", "Giyim", 5, 35),
            ("Bisküvi", "Gıda", 90, 12.75), ("Mont", "Giyim", 12, 900)]
    for r, (u, k, a, f) in enumerate(veri, start=2):
        ws.append([u, k, a, f, f"=C{r}*D{r}"])
    for c in "ABCDE":
        ws[f"{c}1"].font = Font(b=True)
    ws["A11"] = "Görünen toplam"
    ws["E11"] = "=SUBTOTAL(109,E2:E9)"
    ws["A12"] = "Tüm toplam"
    ws["E12"] = "=SUM(E2:E9)"
    ws["G1"] = "KDV"
    ws["G2"] = 0.2
    ws["A9"].comment = Comment("Sezon sonu: fiyat güncellenecek", "Ata")
    ws.freeze_panes = "A2"
    # filtre: Kategori = Gıda ya da Giyim; gizli satırlar dosyada da hidden
    ws.auto_filter.ref = "A1:E9"
    fc = FilterColumn(colId=1)
    fc.filters = Filters(filter=["Gıda", "Giyim"])
    ws.auto_filter.filterColumn.append(fc)
    for r in (2, 3):
        ws.row_dimensions[r].hidden = True
    dv = DataValidation(type="list", formula1='"Elektronik,Gıda,Giyim"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Kategori", error="Listeden seçin")
    ws.add_data_validation(dv)
    dv.add("B2:B9")
    dv2 = DataValidation(type="whole", operator="between", formula1="0", formula2="10000", showErrorMessage=True)
    ws.add_data_validation(dv2)
    dv2.add("C2:C9")
    ws.conditional_formatting.add("C2:C9", CellIsRule(operator="lessThan", formula=["10"],
                                                      fill=PatternFill(bgColor="FFC7CE"), font=Font(color="9C0006")))
    ws.conditional_formatting.add("E2:E9", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    ws.conditional_formatting.add("D2:D9", ColorScaleRule(start_type="min", start_color="63BE7B", end_type="max",
                                                          end_color="F8696B"))
    fy = wb.create_sheet("Fiyatlar")
    fy.append(["Ürün", "Liste fiyatı"])
    for u, _, _, f in veri:
        fy.append([u, f])
    wb.defined_names["FiyatListesi"] = DefinedName("FiyatListesi", attr_text="Fiyatlar!$A$2:$B$9")
    wb.defined_names["KdvOrani"] = DefinedName("KdvOrani", attr_text="Stok!$G$2")
    ws["F1"] = "KDV'li tutar"
    for r in range(2, 10):
        ws[f"F{r}"] = f"=E{r}*(1+KdvOrani)"
    ws["H1"] = "Liste fiyatı"
    for r in range(2, 10):
        ws[f"H{r}"] = f"=VLOOKUP(A{r},FiyatListesi,2,FALSE)"
    wb.save(yol)


def grafikler(yol):
    """Tabcel'in KENDİ yazdığı dosya: 7 tür × alt tür = 13 grafik (galeri) + resim (Faz 6 / 6j). Excel'de
    `araclar\\excel_dogrula.ps1` ile açılıp doğrulanır; testler Tabcel imini (uyarısız açılış) denetler."""
    sys.path.insert(0, os.path.dirname(KLASOR.rsplit(os.sep, 1)[0]))
    from tabcel.io import xlsx_io
    from tabcel.model.adres import Aralik
    from tabcel.model.grafik import Konum, Resim
    from tabcel.model.grafik_islem import GALERI, grafik_kur
    from tabcel.model.sayfa import CalismaKitabi, Hucre
    import io
    from PIL import Image as PilImage

    k = CalismaKitabi()
    s = k.sayfalar[0]
    for c, b in enumerate(["Ay", "Satış", "Maliyet", "Kâr"], start=1):
        s.ata(1, c, Hucre(b))
    for i, (ay, a, m) in enumerate([("Ocak", 120, 80), ("Şubat", 135, 90), ("Mart", 150, 95), ("Nisan", 142, 100),
                                    ("Mayıs", 160, 110), ("Haziran", 175, 120)]):
        r = i + 2
        s.ata(r, 1, Hucre(ay))
        s.ata(r, 2, Hucre(float(a)))
        s.ata(r, 3, Hucre(float(m)))
        s.ata(r, 4, Hucre(formul=f"=B{r}-C{r}"))
    sira = 0
    for tur, altlar in GALERI.items():
        for alt, ad in altlar:
            g = grafik_kur(s, Aralik(1, 1, 7, 3), tur, alt)
            g = g.degistir(baslik=ad, konum=Konum(2 + (sira // 2) * 16, 6 + (sira % 2) * 9), boyut=(480, 288))
            s.nesne_ekle(g)
            sira += 1
    b = io.BytesIO()
    PilImage.new("RGB", (160, 80), (33, 115, 70)).save(b, "PNG")
    s.nesne_ekle(Resim(veri=b.getvalue(), ad="logo", konum=Konum(10, 1), boyut=(160, 80)))
    xlsx_io.kaydet(k, yol)


def csvler():
    with open(os.path.join(KLASOR, "satis_tr.csv"), "wb") as f:
        f.write("Ürün;Adet;Fiyat;Tarih\r\nÇay;10;12,5;12.10.2026\r\nŞeker;4;8,25;13.10.2026\r\n"
                "İncir;7;40;14.10.2026\r\n".encode("cp1254"))
    with open(os.path.join(KLASOR, "satis_utf8.csv"), "wb") as f:
        f.write('Ad,Not\r\n"Ali, Veli",merhaba\r\nAyşe,"iki\r\nsatır"\r\n'.encode("utf-8"))


def main():
    os.makedirs(KLASOR, exist_ok=True)
    butce(os.path.join(KLASOR, "butce.xlsx"))
    grafikli(os.path.join(KLASOR, "grafikli.xlsx"))
    stok(os.path.join(KLASOR, "stok.xlsx"))
    grafikler(os.path.join(KLASOR, "grafikler.xlsx"))
    csvler()
    print("Üretildi:", ", ".join(sorted(os.listdir(KLASOR))))


if __name__ == "__main__":
    sys.exit(main())
