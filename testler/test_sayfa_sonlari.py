"""T9: el ile sayfa sonları, sayfa sınırları, yüksekliğe sığdır, sayfa sırası (model + .xlsx; Qt'siz)."""
import pytest

from tabcel.io import xlsx_io
from tabcel.model.adres import Aralik
from tabcel.model.komutlar import GeriAlYigini, SatirEkleSil, SayfaYapisiDegis, SutunEkleSil
from tabcel.model.sayfa import CalismaKitabi, Hucre
from tabcel.model.sayfa_yapisi import SayfaYapisi
from tabcel.model.sayfalama import sayfa_sinirlari, sayfala


def _s(satir=100, sutun=3):
    s = CalismaKitabi().sayfalar[0]
    for r in range(1, satir + 1):
        for c in range(1, sutun + 1):
            s.ata(r, c, Hucre(float(r)))
    return s


def _olcu(s):
    return (lambda c: s.genislik(c)), (lambda r: s.yukseklik(r))


# ---------------------------------------------------------------- doğrulama
def test_kesmeler_siralanir_tekrarsiz():
    assert SayfaYapisi(satir_kesmeleri=[30, 10, 10]).satir_kesmeleri == (10, 30)


@pytest.mark.parametrize("kw", [{"satir_kesmeleri": (1,)}, {"sutun_kesmeleri": (0,)},
                                {"satir_kesmeleri": (1048577,)}, {"sutun_kesmeleri": (16385,)},
                                {"sayfa_sirasi": "x"}])
def test_gecersiz_kesme_ve_sira(kw):
    with pytest.raises(ValueError):
        SayfaYapisi(**kw)


# ---------------------------------------------------------------- bölme
def test_el_ile_satir_sonu_yeni_sayfa_baslatir():
    s = _s(30)
    y = SayfaYapisi(satir_kesmeleri=(11, 21))
    p = sayfala(s, y, *_olcu(s))
    assert [d.satirlar[0] for d in p] == [1, 11, 21]
    assert [d.satirlar[-1] for d in p] == [10, 20, 30]


def test_el_ile_sutun_sonu():
    s = _s(5, 6)
    p = sayfala(s, SayfaYapisi(sutun_kesmeleri=(3, 5)), *_olcu(s))
    assert [d.sutunlar for d in p] == [(1, 2), (3, 4), (5, 6)]


def test_otomatik_bolme_kesmeden_sonra_surer():
    s = _s(300)
    otomatik = sayfala(s, SayfaYapisi(), *_olcu(s))
    ilk_uzunluk = len(otomatik[0].satirlar)
    y = SayfaYapisi(satir_kesmeleri=(5,))
    p = sayfala(s, y, *_olcu(s))
    assert p[0].satirlar == (1, 2, 3, 4) and p[1].satirlar[0] == 5
    assert len(p[1].satirlar) == ilk_uzunluk                       # kesmeden sonra otomatik sayfa dolana dek sürer


def test_kesme_baslik_satirini_sonraki_sayfaya_yineler():
    s = _s(30)
    y = SayfaYapisi(satir_kesmeleri=(11,), baslik_satirlari=(1, 1))
    p = sayfala(s, y, *_olcu(s))
    assert p[1].satirlar[:2] == (1, 11) and p[1].baslik_sayisi == 1 and p[0].baslik_sayisi == 0


@pytest.mark.parametrize("kw", [{"sigdir_genislik": True}, {"sigdir_yukseklik": True}])
def test_sigdir_acikken_kesmeler_yok_sayilir(kw):
    s = _s(30)
    p = sayfala(s, SayfaYapisi(satir_kesmeleri=(11,), **kw), *_olcu(s))
    assert len(p) == 1
    sn = sayfa_sinirlari(s, SayfaYapisi(satir_kesmeleri=(11,), **kw), *_olcu(s))
    assert sn.satirlar == ()


def test_yukseklige_sigdir_olcegi_kucultur():
    s = _s(300)
    tek = sayfala(s, SayfaYapisi(sigdir_yukseklik=True), *_olcu(s))
    assert len(tek) == 1 and 0.1 <= tek[0].olcek < 1.0
    assert len(sayfala(s, SayfaYapisi(), *_olcu(s))) > 1
    yy = SayfaYapisi()
    ikisi = sayfala(s, yy.degistir(sigdir_yukseklik=True, sigdir_genislik=True), *_olcu(s))
    assert ikisi[0].olcek <= tek[0].olcek + 1e-9                    # ikisi birden daha küçük ya da eşit


def test_sayfa_sirasi_numaralamayi_degistirir():
    s = _s(100, 40)
    y = SayfaYapisi(satir_kesmeleri=(11,), sutun_kesmeleri=(5,))
    asagi = sayfala(s, y, *_olcu(s))
    sag = sayfala(s, y.degistir(sayfa_sirasi="sag_sonra_asagi"), *_olcu(s))
    assert len(asagi) == len(sag) and len(asagi) >= 4
    # önce aşağı: 2. sayfa aynı sütun diliminin alt parçası; önce sağa: 2. sayfa sağ komşu
    assert asagi[1].sutunlar == asagi[0].sutunlar and asagi[1].satirlar != asagi[0].satirlar
    assert sag[1].satirlar == sag[0].satirlar and sag[1].sutunlar != sag[0].sutunlar
    assert {(d.satirlar, d.sutunlar) for d in asagi} == {(d.satirlar, d.sutunlar) for d in sag}


# ---------------------------------------------------------------- ekran sınırları
def test_sayfa_sinirlari_sayfala_ile_uyumlu():
    s = _s(300, 30)
    for y in (SayfaYapisi(), SayfaYapisi(satir_kesmeleri=(7, 50), sutun_kesmeleri=(9,)),
              SayfaYapisi(baslik_satirlari=(1, 2)), SayfaYapisi(olcek=50)):
        d = sayfala(s, y, *_olcu(s))
        sn = sayfa_sinirlari(s, y, *_olcu(s))
        satir_baslar = sorted({x.satirlar[x.baslik_sayisi] for x in d})[1:]
        sutun_baslar = sorted({x.sutunlar[0] for x in d})[1:]
        assert [r for r, _ in sn.satirlar] == satir_baslar
        assert [c for c, _ in sn.sutunlar] == sutun_baslar


def test_sinirlarda_el_ile_ve_otomatik_ayrilir():
    s = _s(200)
    sn = sayfa_sinirlari(s, SayfaYapisi(satir_kesmeleri=(5,)), *_olcu(s))
    assert sn.satirlar[0] == (5, True)
    assert all(not el for _, el in sn.satirlar[1:]) and len(sn.satirlar) >= 2
    assert sn.alan == Aralik(1, 1, 200, 3)


def test_bos_sayfada_sinir_yok():
    s = CalismaKitabi().sayfalar[0]
    assert sayfa_sinirlari(s, SayfaYapisi(), *_olcu(s)) is None


# ---------------------------------------------------------------- ekle-sil kayması
def test_satir_ekle_sil_kesmeleri_kaydirir_ve_geri_alir():
    s = _s(40)
    s.sayfa_yapisi = SayfaYapisi(satir_kesmeleri=(10, 20), alan=Aralik(1, 1, 30, 3), baslik_satirlari=(1, 2))
    y = GeriAlYigini()
    y.calistir(SatirEkleSil(s, 5, 2, True))                         # 5'e 2 satır ekle: hepsi 2 aşağı
    assert s.sayfa_yapisi.satir_kesmeleri == (12, 22) and s.sayfa_yapisi.alan == Aralik(1, 1, 32, 3)
    assert s.sayfa_yapisi.baslik_satirlari == (1, 2)
    y.geri_al()
    assert s.sayfa_yapisi.satir_kesmeleri == (10, 20) and s.sayfa_yapisi.alan == Aralik(1, 1, 30, 3)
    y.calistir(SatirEkleSil(s, 8, 5, False))                        # 8..12 silinir: 10 düşer, 20 → 15
    assert s.sayfa_yapisi.satir_kesmeleri == (15,) and s.sayfa_yapisi.alan == Aralik(1, 1, 25, 3)
    y.geri_al()
    assert s.sayfa_yapisi.satir_kesmeleri == (10, 20)


def test_baslik_satirlari_satir_eklenince_kayar():
    s = _s(10)
    s.sayfa_yapisi = SayfaYapisi(baslik_satirlari=(3, 4))
    y = GeriAlYigini()
    y.calistir(SatirEkleSil(s, 1, 1, True))
    assert s.sayfa_yapisi.baslik_satirlari == (4, 5)
    y.calistir(SatirEkleSil(s, 5, 1, True))                         # başlığın içine ekleme: büyür
    assert s.sayfa_yapisi.baslik_satirlari == (4, 6)


def test_sutun_ekle_sil_sutun_kesmelerini_kaydirir():
    s = _s(5, 10)
    s.sayfa_yapisi = SayfaYapisi(sutun_kesmeleri=(4, 8), satir_kesmeleri=(3,))
    y = GeriAlYigini()
    y.calistir(SutunEkleSil(s, 1, 1, True))
    assert s.sayfa_yapisi.sutun_kesmeleri == (5, 9) and s.sayfa_yapisi.satir_kesmeleri == (3,)
    y.geri_al()
    assert s.sayfa_yapisi.sutun_kesmeleri == (4, 8)


def test_sayfa_yapisi_degis_kesmeleri_geri_alir():
    s = _s(10)
    y = GeriAlYigini()
    y.calistir(SayfaYapisiDegis(s, s.sayfa_yapisi.degistir(satir_kesmeleri=(5,))))
    assert s.sayfa_yapisi.satir_kesmeleri == (5,)
    y.geri_al()
    assert s.sayfa_yapisi.satir_kesmeleri == ()
    y.yinele()
    assert s.sayfa_yapisi.satir_kesmeleri == (5,)


# ---------------------------------------------------------------- .xlsx
def _gidis_donus(yapi, tmp_path):
    k = CalismaKitabi()
    k.sayfalar[0].ata(1, 1, Hucre("x"))
    k.sayfalar[0].sayfa_yapisi = yapi
    xlsx_io.kaydet(k, str(tmp_path / "a.xlsx"))
    k2, rapor = xlsx_io.ac(str(tmp_path / "a.xlsx"))
    return k2.sayfalar[0].sayfa_yapisi, rapor


@pytest.mark.parametrize("yapi", [
    SayfaYapisi(satir_kesmeleri=(11, 30), sutun_kesmeleri=(4,)),
    SayfaYapisi(sayfa_sirasi="sag_sonra_asagi"),
    SayfaYapisi(sigdir_yukseklik=True),
    SayfaYapisi(sigdir_genislik=True),
    SayfaYapisi(sigdir_genislik=True, sigdir_yukseklik=True),
    SayfaYapisi(olcek=80, satir_kesmeleri=(2,)),
])
def test_xlsx_gidis_donus(yapi, tmp_path):
    geri, rapor = _gidis_donus(yapi, tmp_path)
    assert geri == yapi and not rapor.veri_kaybi


def test_xlsx_excel_kesme_kimligi_onceki_satirdir(tmp_path):
    """Excel `id=10` = 10. satırdan SONRA kesme → yeni sayfa 11'de başlar (Tabcel'de 11)."""
    from openpyxl import Workbook
    from openpyxl.worksheet.pagebreak import Break
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.row_breaks.append(Break(id=10, man=True))
    ws.row_breaks.append(Break(id=20, man=False))                   # otomatik: okunmaz
    ws.col_breaks.append(Break(id=3, man=True))
    wb.save(tmp_path / "e.xlsx")
    k, _ = xlsx_io.ac(str(tmp_path / "e.xlsx"))
    y = k.sayfalar[0].sayfa_yapisi
    assert y.satir_kesmeleri == (11,) and y.sutun_kesmeleri == (4,)


def test_xlsx_baska_programin_sigdir_dosyasi(tmp_path):
    """fitToWidth/Height yazılmamışsa 1 sayılır (Excel); 0 = serbest."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(tmp_path / "f.xlsx")
    y = xlsx_io.ac(str(tmp_path / "f.xlsx"))[0].sayfalar[0].sayfa_yapisi
    assert y.sigdir_genislik and y.sigdir_yukseklik
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    wb.save(tmp_path / "g.xlsx")
    y = xlsx_io.ac(str(tmp_path / "g.xlsx"))[0].sayfalar[0].sayfa_yapisi
    assert y.sigdir_genislik and not y.sigdir_yukseklik
