"""Hızlı hücre okuyucu (`io/xlsx_hizli.py`): özgün openpyxl okuyucusuyla **birebir aynı kitap**. Tüm örnek dosyalar + özel durumlar
(formül, paylaşılan/dizi formül, satır içi metin, `&` varlıkları, hata, mantıksal, tarih/saat/süre, 1904, birleşim, köprü, yorum, `_tabcel`)."""
import glob
import os
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from tabcel.io import xlsx_hizli, xlsx_io
from tabcel.model.sayfa import CalismaKitabi, Hucre

ORNEKLER = os.path.join(os.path.dirname(__file__), "ornekler")


def dokum(k: CalismaKitabi):
    """Karşılaştırılacak her şey: hücreler, biçimler (Stil nesnesi), boyutlar, birleşimler, notlar, doğrulama/koşullu/filtre sayıları, adlar."""
    cikti = []
    for s in k.sayfalar:
        cikti.append((
            s.ad, s.gizli, getattr(s, "cok_gizli", False),
            sorted(((a, (h.deger, h.formul)) for a, h in s._hucreler.items()), key=lambda x: x[0]),
            sorted(((a, repr(k.stil_havuzu.stil(v))) for a, v in s._stil.items()), key=lambda x: x[0]),
            sorted((a, repr(k.stil_havuzu.stil(v))) for a, v in s.satir_stil.items()),
            sorted((a, repr(k.stil_havuzu.stil(v))) for a, v in s.sutun_stil.items()),
            sorted(s.sutun_genislik.items()), sorted(s.satir_yukseklik.items()), sorted(s.gizli_satir), sorted(s.gizli_sutun),
            [repr(b) for b in s.birlesimler], sorted((a, repr(n)) for a, n in s.notlar.items()),
            len(s.dogrulamalar), len(s.kosullu), repr(s.filtre), repr(s.sayfa_yapisi), len(s.nesneler),
        ))
    cikti.append(sorted(repr(a) for a in k.adlar.values()) if hasattr(k, "adlar") else None)
    return cikti


def ac(yol, hizli: bool):
    (xlsx_hizli.kur if hizli else xlsx_hizli.kapat)()
    try:
        k, rapor = xlsx_io.ac(yol)
    finally:
        xlsx_hizli.kur()
    return k, rapor


def esit_mi(yol):
    k1, r1 = ac(yol, False)
    k2, r2 = ac(yol, True)
    assert dokum(k1) == dokum(k2), os.path.basename(str(yol))
    assert sorted(r1.uyarilar) == sorted(r2.uyarilar)
    return k2


# ---------------------------------------------------------------- tüm örnek dosyalar
@pytest.mark.parametrize("yol", sorted(glob.glob(os.path.join(ORNEKLER, "*.xlsx")) + glob.glob(os.path.join(ORNEKLER, "excel", "*.xlsx"))),
                         ids=lambda y: os.path.relpath(y, ORNEKLER))
def test_ornek_dosyalar_ozgunle_ayni(yol):
    esit_mi(yol)


def test_yama_varsayilan_olarak_acik():
    assert xlsx_hizli.yamali_mi()


# ---------------------------------------------------------------- özel durumlar (openpyxl ile üretilen dosyalar)
def _yaz(tmp_path, ad, doldur, **kw):
    wb = Workbook(**kw)
    ws = wb.active
    doldur(ws, wb)
    yol = tmp_path / ad
    wb.save(yol)
    return str(yol)


def test_deger_turleri_ve_kacis(tmp_path):
    import datetime

    def d(ws, wb):
        ws["A1"] = "Ürün & <Ad> \"x\""
        ws["A2"] = 12
        ws["A3"] = 1234.5
        ws["A4"] = True
        ws["A5"] = False
        ws["A6"] = "#DIV/0!"
        ws["A6"].data_type = "e"
        ws["A7"] = datetime.datetime(2026, 10, 15, 13, 30, 15)
        ws["A8"] = datetime.date(2026, 1, 1)
        ws["A9"] = datetime.time(6, 0)
        ws["A10"] = datetime.timedelta(hours=30)
        ws["A11"] = "satır\nsonu"
        ws["A12"] = "  baştaki boşluk  "
        ws["A13"] = ""
        ws["A14"] = "007"
        ws["A15"] = 1e-10
        ws["A16"] = 123456789012345678
    k = esit_mi(_yaz(tmp_path, "t.xlsx", d))
    s = k.sayfalar[0]
    assert s.deger(1, 1) == "Ürün & <Ad> \"x\"" and s.deger(4, 1) is True and s.deger(11, 1) == "satır\nsonu"
    assert s.deger(7, 1) == pytest.approx(46310.5625 + 15 / 86400, rel=1e-9)


def test_formuller_paylasilan_dizi_ve_bos_formul(tmp_path):
    def d(ws, wb):
        for r in range(1, 6):
            ws.cell(r, 1, r)
            ws.cell(r, 2, f"=A{r}*2")
        ws["C1"] = "=SUM(A1:A5)"
        ws["D1"] = '=IF(A1>0,"evet","hayır")'
        ws["E1"] = "=A1&\"&\""
        ws.formula_attributes = {}
    k = esit_mi(_yaz(tmp_path, "f.xlsx", d))
    assert k.sayfalar[0].hucre(3, 2).formul == "=A3*2"


def test_yorum_koprü_birlesim_stil(tmp_path):
    def d(ws, wb):
        ws["A1"] = "başlık"
        ws["A1"].font = Font(bold=True, color="FF0000")
        ws["A1"].fill = PatternFill("solid", start_color="FFFF00")
        ws["B1"] = "bağlantı"
        ws["B1"].hyperlink = "https://ornek.com/yol"
        ws["C1"] = "yorumlu"
        ws["C1"].comment = Comment("not metni", "yazar")
        ws["A3"] = "birleşik"
        ws.merge_cells("A3:C4")
        ws["E1"].font = Font(italic=True)                   # değeri yok, yalnız biçim
        ws.row_dimensions[2].height = 33
        ws.column_dimensions["B"].width = 25
        ws.row_dimensions[5].hidden = True
        ws["A5"] = 1
    k = esit_mi(_yaz(tmp_path, "y.xlsx", d))
    s = k.sayfalar[0]
    assert s.deger(1, 2) == "bağlantı" and s.deger(1, 3) == "yorumlu" and (1, 3) in s.notlar


def test_1904_tarih_sistemi_ozgun_yola_duser(tmp_path):
    import datetime

    def d(ws, wb):
        wb.epoch = __import__("openpyxl.utils.datetime", fromlist=["CALENDAR_MAC_1904"]).CALENDAR_MAC_1904
        ws["A1"] = datetime.date(2026, 1, 1)
        ws["A2"] = 5
    esit_mi(_yaz(tmp_path, "e.xlsx", d))


def test_gizli_sayfalar_ve_cok_sayfa(tmp_path):
    def d(ws, wb):
        ws["A1"] = "bir"
        o = wb.create_sheet("İkinci")
        o["B2"] = "iki"
        g = wb.create_sheet("Gizli")
        g["A1"] = "g"
        g.sheet_state = "veryHidden"
    esit_mi(_yaz(tmp_path, "g.xlsx", d))


def test_bos_sayfa_ve_tek_hucre(tmp_path):
    esit_mi(_yaz(tmp_path, "b.xlsx", lambda ws, wb: None))
    esit_mi(_yaz(tmp_path, "t.xlsx", lambda ws, wb: ws.__setitem__("Z100", "x")))


def test_satir_icinde_metin_write_only(tmp_path):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Veri")
    ws.append(["Ad", "Sayı", "Not"])
    for i in range(300):
        ws.append([f"Ürün {i} & ç", i, None if i % 3 else "x"])
    yol = tmp_path / "w.xlsx"
    wb.save(yol)
    k = esit_mi(str(yol))
    assert k.sayfalar[0].deger(301, 1) == "Ürün 299 & ç"


def test_tabcel_ile_kaydedilen_kitap_ayni_okunur(tmp_path):
    k = CalismaKitabi()
    s = k.sayfalar[0]
    for r in range(1, 40):
        s.ata(r, 1, Hucre(f"Satır {r}"))
        s.ata(r, 2, Hucre(float(r)))
        s.ata(r, 3, Hucre(None, f"=B{r}*2"))
    yol = str(tmp_path / "tabcel.xlsx")
    xlsx_io.kaydet(k, yol)
    esit_mi(yol)


def test_sablon_bilgisi_hizli_okuyucuda_korunur(tmp_path):
    """`_tabcel` çok gizli sayfası `iter_rows` ile okunur: bu sayfa Cell nesnesi olarak yüklenmeli."""
    from pathlib import Path
    from tabcel.sablon import yukleyici
    from tabcel.sablon.uretici import uret
    d = yukleyici.depo_yukle(yukleyici.yerlesik_klasor(), Path("yok_klasor"))
    r = uret(d.bul("yerlesik:stok_takibi"), None, ornek_veri=True)
    yol = str(tmp_path / "s.xlsx")
    xlsx_io.kaydet(r.kitap, yol)
    k = esit_mi(yol)
    assert k.sablon is not None and k.sablon_denetim.aktif


def test_bozuk_xml_ozgun_hata_verir(tmp_path):
    yol = str(tmp_path / "bozuk.xlsx")
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.save(yol)
    with zipfile.ZipFile(yol) as z:
        girdiler = {n: z.read(n) for n in z.namelist()}
    girdiler["xl/worksheets/sheet1.xml"] = girdiler["xl/worksheets/sheet1.xml"].replace(b"</row>", b"<row")
    with zipfile.ZipFile(yol, "w") as z:
        for n, v in girdiler.items():
            z.writestr(n, v)
    with pytest.raises(xlsx_io.XlsxHatasi):
        xlsx_io.ac(yol)


def test_ortam_degiskeni_yamayi_kapatir(monkeypatch):
    xlsx_hizli.kapat()
    monkeypatch.setenv("TABCEL_YAVAS_OKUYUCU", "1")
    assert xlsx_hizli.kur() is False
    monkeypatch.delenv("TABCEL_YAVAS_OKUYUCU")
    assert xlsx_hizli.kur() is True


def test_hizli_okuyucu_gercekten_hizli(tmp_path):
    """Gerilemeyi yakalar: 40.000 hücrelik dosyada hızlı yol özgün yoldan en az 1,8× hızlı (gevşek eşik; ölçüm ≈3×)."""
    import time
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("V")
    for i in range(5000):
        ws.append([f"Ürün {i % 97}", i, i * 1.5, "x", None, i % 2 == 0, f"n{i}", 5])
    yol = tmp_path / "h.xlsx"
    wb.save(yol)
    sure = {}
    for ad, hizli in (("ozgun", False), ("hizli", True)):
        en_iyi = 1e9
        for _ in range(2):
            t = time.perf_counter()
            ac(str(yol), hizli)
            en_iyi = min(en_iyi, time.perf_counter() - t)
        sure[ad] = en_iyi
    assert sure["hizli"] * 1.8 < sure["ozgun"], sure


def test_ac_disinda_load_workbook_ozgun_anlami_korur(tmp_path):
    """Demet modu yalnız `xlsx_io.ac` içinde açık: başka `load_workbook` çağrıları (testler, modüller) normal `Cell` görür."""
    from openpyxl import load_workbook
    yol = _yaz(tmp_path, "d.xlsx", lambda ws, wb: (ws.__setitem__("A1", 7), ws.__setitem__("B2", "x")))
    ws = load_workbook(yol).active
    assert ws["A1"].value == 7 and ws["B2"].value == "x" and len(ws._cells) == 2
