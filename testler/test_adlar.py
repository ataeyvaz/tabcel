"""Faz 5 / 5d: tanımlı ad (adlandırılmış aralık) — doğrulama, değerlendirme, motor, komut, ekle-sil, taşı.
Qt yok. Kurallar: faz5_tasarim.md §4."""
import pytest

from tabcel.formul import ceviri
from tabcel.formul.kaydir import adi_degistir
from tabcel.formul.motor import Hesaplayici
from tabcel.model import deger as D
from tabcel.model.ad import AdHatasi, TanimliAd, ad_anahtari, ad_dogrula
from tabcel.model.komutlar import AdDegis, GeriAlYigini, HucreTasi, SatirEkleSil, SayfaSil
from tabcel.model.sayfa import CalismaKitabi, Hucre


def kur():
    k = CalismaKitabi()
    s = k.sayfalar[0]
    m = Hesaplayici(k)
    return k, s, m


def f(k, metin):
    return Hucre(None, ceviri.girdiden(metin, k))


def ekle(k, ad, formul, kapsam=None):
    k.ad_yaz(TanimliAd(ad, formul, kapsam))


class TestDogrulama:
    @pytest.mark.parametrize("ad", ["Fiyat", "KDV.Oranı", "_x", "Çeyrek_1", "\\x", "XFE1", "a1b", "Rate"])
    def test_gecerli(self, ad):
        ad_dogrula(ad)

    @pytest.mark.parametrize("ad", ["", "A1", "XFD1048576", "KDV2", "R", "C", "r", "c", "R1C1", "R5", "C5", "RC",
                                    "Toplam 1", "1x", "a-b", "DOĞRU", "YANLIŞ", "true", "False", "x" * 256, " x"])
    def test_gecersiz(self, ad):
        with pytest.raises(AdHatasi):
            ad_dogrula(ad)

    def test_255_karakter_gecerli(self):
        ad_dogrula("x" * 255)

    def test_kontrol_yinelenen_buyuk_kucuk_duyarsiz(self):
        k, s, m = kur()
        ekle(k, "Fiyat", "=[1]!$A$1")
        with pytest.raises(AdHatasi):
            AdDegis(k, None, TanimliAd("FİYAT".replace("İ", "i"), "=1"))
        AdDegis(k, None, TanimliAd("Fiyat", "=1", kapsam=s.kimlik))        # başka kapsam serbest

    def test_kontrol_kapsam_yok_ve_formul_yok(self):
        k, s, m = kur()
        with pytest.raises(AdHatasi):
            AdDegis(k, None, TanimliAd("X", "=1", kapsam=99))
        with pytest.raises(AdHatasi):
            AdDegis(k, None, TanimliAd("X", "A1"))


class TestDegerlendirme:
    def test_aralik_adi(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): Hucre(10.0), (2, 1): Hucre(20.0), (3, 1): f(k, "=TOPLA(Fiyat)")})
        ekle(k, "Fiyat", "=[1]!$A$1:$A$2")
        assert s.sonuc(3, 1) == 30.0

    def test_sabit_ad(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): Hucre(100.0), (2, 1): f(k, "=A1*KDV")})
        ekle(k, "KDV", "=0.18")
        assert s.sonuc(2, 1) == pytest.approx(18.0)

    def test_buyuk_kucuk_harf_duyarsiz(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): f(k, "=kdv+1")})
        ekle(k, "KDV", "=2")
        assert s.sonuc(1, 1) == 3.0

    def test_turkce_buyuk_i(self):
        k, s, m = kur()
        ekle(k, "Işık", "=5")
        s.toplu_yaz({(1, 1): f(k, "=IŞIK*2")})
        assert s.sonuc(1, 1) == 10.0

    def test_tanimsiz_ad_hata(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): f(k, "=Yok+1")})
        assert s.sonuc(1, 1) == D.AD

    def test_ad_sonradan_tanimlaninca_hesaplanir(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): f(k, "=Yok+1")})
        ekle(k, "Yok", "=4")
        assert s.sonuc(1, 1) == 5.0

    def test_silinen_ad_hata_formul_metni_korunur(self):
        k, s, m = kur()
        ekle(k, "KDV", "=2")
        s.toplu_yaz({(1, 1): f(k, "=KDV+1")})
        assert s.sonuc(1, 1) == 3.0
        k.ad_yaz(None, ad_anahtari(None, "KDV"))
        assert s.sonuc(1, 1) == D.AD
        assert s.hucre(1, 1).formul == "=KDV+1"

    def test_degisen_hucre_yeniden_hesaplanir(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): Hucre(1.0), (2, 1): Hucre(2.0), (3, 1): f(k, "=TOPLA(Fiyat)")})
        ekle(k, "Fiyat", "=[1]!$A$1:$A$2")
        s.toplu_yaz({(2, 1): Hucre(50.0)})
        assert s.sonuc(3, 1) == 51.0

    def test_ad_formulu_degisince_yeniden_hesaplanir(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): Hucre(1.0), (2, 1): Hucre(7.0), (3, 1): f(k, "=Hedef*2")})
        ekle(k, "Hedef", "=[1]!$A$1")
        assert s.sonuc(3, 1) == 2.0
        ekle(k, "Hedef", "=[1]!$A$2")
        assert s.sonuc(3, 1) == 14.0

    def test_baska_sayfa_adi(self):
        k, s, m = kur()
        s2 = k.sayfa_ekle("Veri")
        s2.toplu_yaz({(1, 1): Hucre(3.0), (2, 1): Hucre(4.0)})
        ekle(k, "Liste", f"=[{s2.kimlik}]!$A$1:$A$2")
        s.toplu_yaz({(1, 1): f(k, "=TOPLA(Liste)")})
        assert s.sonuc(1, 1) == 7.0
        s2.toplu_yaz({(2, 1): Hucre(10.0)})
        assert s.sonuc(1, 1) == 13.0

    def test_ad_baska_adi_kullanir(self):
        k, s, m = kur()
        ekle(k, "Taban", "=10")
        ekle(k, "Fiyat", "=Taban*2")
        s.toplu_yaz({(1, 1): f(k, "=Fiyat+1")})
        assert s.sonuc(1, 1) == 21.0
        ekle(k, "Taban", "=1")
        assert s.sonuc(1, 1) == 3.0

    def test_ozyinelemeli_ad(self):
        k, s, m = kur()
        ekle(k, "A", "=B+1")
        ekle(k, "B", "=A+1")
        s.toplu_yaz({(1, 1): f(k, "=A")})
        assert s.sonuc(1, 1) == D.AD

    def test_kendine_basvuran_ad(self):
        k, s, m = kur()
        ekle(k, "Dongu", "=Dongu+1")
        s.toplu_yaz({(1, 1): f(k, "=Dongu")})
        assert s.sonuc(1, 1) == D.AD

    def test_hucre_dongusu_ad_uzerinden(self):
        k, s, m = kur()
        ekle(k, "Bu", "=[1]!$A$1")
        s.toplu_yaz({(1, 1): f(k, "=Bu+1")})
        assert (s.kimlik, 1, 1) in m.donguler

    def test_hata_degeri_ad(self):
        k, s, m = kur()
        ekle(k, "Bozuk", "=1/0")
        s.toplu_yaz({(1, 1): f(k, "=Bozuk")})
        assert s.sonuc(1, 1) == D.BOL0

    def test_metin_ve_bos(self):
        k, s, m = kur()
        ekle(k, "Selam", '="merhaba"')
        ekle(k, "BosHucre", "=[1]!$Z$9")
        s.toplu_yaz({(1, 1): f(k, "=Selam&\"!\""), (2, 1): f(k, "=BosHucre")})
        assert s.sonuc(1, 1) == "merhaba!"
        assert s.sonuc(2, 1) == 0.0

    def test_bozuk_ad_formulu(self):
        k, s, m = kur()
        ekle(k, "Bozuk", "=1+")
        s.toplu_yaz({(1, 1): f(k, "=Bozuk")})
        assert s.sonuc(1, 1) == D.AD


class TestKapsam:
    def test_sayfa_kapsami_kitabi_golgeler(self):
        k, s, m = kur()
        s2 = k.sayfa_ekle("S2")
        ekle(k, "Oran", "=1")
        ekle(k, "Oran", "=2", kapsam=s2.kimlik)
        s.toplu_yaz({(1, 1): f(k, "=Oran")})
        s2.toplu_yaz({(1, 1): f(k, "=Oran")})
        assert s.sonuc(1, 1) == 1.0
        assert s2.sonuc(1, 1) == 2.0

    def test_sayfa_kapsami_baska_sayfada_gorunmez(self):
        k, s, m = kur()
        s2 = k.sayfa_ekle("S2")
        ekle(k, "Yerel", "=5", kapsam=s2.kimlik)
        s.toplu_yaz({(1, 1): f(k, "=Yerel")})
        s2.toplu_yaz({(1, 1): f(k, "=Yerel")})
        assert s.sonuc(1, 1) == D.AD
        assert s2.sonuc(1, 1) == 5.0

    def test_ayni_ad_iki_sayfada(self):
        k, s, m = kur()
        s2 = k.sayfa_ekle("S2")
        ekle(k, "Y", "=1", kapsam=s.kimlik)
        ekle(k, "Y", "=2", kapsam=s2.kimlik)
        s.toplu_yaz({(1, 1): f(k, "=Y")})
        s2.toplu_yaz({(1, 1): f(k, "=Y")})
        assert (s.sonuc(1, 1), s2.sonuc(1, 1)) == (1.0, 2.0)

    def test_sayfa_silinince_ad_kalir_geri_alinca_doner(self):
        k, s, m = kur()
        s2 = k.sayfa_ekle("S2")
        ekle(k, "Yerel", "=5", kapsam=s2.kimlik)
        s2.toplu_yaz({(1, 1): f(k, "=Yerel")})
        y = GeriAlYigini()
        y.calistir(SayfaSil(k, s2))
        assert ad_anahtari(s2.kimlik, "Yerel") in k.adlar          # silinmez
        assert k.ad_bul("Yerel", s2.kimlik) is not None            # kapsam kimliği ile hâlâ bulunur
        y.geri_al()
        assert s2.sonuc(1, 1) == 5.0


class TestGoreli:
    def test_goreli_ad_a1e_gore(self):
        """Ad A1'e göre saklanır: `=[1]!A1` (göreli) hücrenin kendi konumundan aynı kayma uygulanır."""
        k, s, m = kur()
        ekle(k, "SolHucre", "=[1]!A1")           # A1'e göre: bulunduğu hücrenin KENDİSİ
        ekle(k, "Sag", "=[1]!B1")                # bulunduğu hücrenin sağındaki
        s.toplu_yaz({(1, 3): Hucre(5.0), (2, 2): Hucre(1.0), (2, 3): f(k, "=Sag")})
        # B2... ad hücre C2'de kullanılıyor: Sag = (2,3)+(0,1) = D2 → boş → 0
        assert s.sonuc(2, 3) == 0.0
        s.toplu_yaz({(2, 4): Hucre(9.0)})
        assert s.sonuc(2, 3) == 9.0

    def test_goreli_ad_motor_bagimliligi(self):
        k, s, m = kur()
        ekle(k, "Ustte", "=[1]!A1")
        s.toplu_yaz({(1, 2): Hucre(4.0), (2, 2): f(k, "=Ustte*2")})   # B2: A1 + (1,1) = B2 → kendisi
        # B2'den Ustte = (1,1)+(1,1) = B2 → döngü
        assert (s.kimlik, 2, 2) in m.donguler

    def test_goreli_ad_baska_hucrede(self):
        k, s, m = kur()
        ekle(k, "Solu", "=[1]!$A1")                # $A: sütun sabit, satır göreli
        s.toplu_yaz({(1, 1): Hucre(1.0), (2, 1): Hucre(2.0), (1, 3): f(k, "=Solu"), (2, 3): f(k, "=Solu")})
        assert (s.sonuc(1, 3), s.sonuc(2, 3)) == (1.0, 2.0)


class TestKomut:
    def test_ekle_geri_al_yinele(self):
        k, s, m = kur()
        s.toplu_yaz({(1, 1): f(k, "=KDV+1")})
        y = GeriAlYigini()
        y.calistir(AdDegis(k, None, TanimliAd("KDV", "=2")))
        assert s.sonuc(1, 1) == 3.0
        y.geri_al()
        assert k.adlar == {} and s.sonuc(1, 1) == D.AD
        y.yinele()
        assert s.sonuc(1, 1) == 3.0

    def test_duzenle(self):
        k, s, m = kur()
        ekle(k, "KDV", "=2")
        s.toplu_yaz({(1, 1): f(k, "=KDV+1")})
        y = GeriAlYigini()
        y.calistir(AdDegis(k, ad_anahtari(None, "KDV"), TanimliAd("KDV", "=10", aciklama="oran")))
        assert s.sonuc(1, 1) == 11.0
        assert k.ad_bul("KDV").aciklama == "oran"
        y.geri_al()
        assert s.sonuc(1, 1) == 3.0 and k.ad_bul("KDV").aciklama == ""

    def test_sil(self):
        k, s, m = kur()
        ekle(k, "KDV", "=2")
        s.toplu_yaz({(1, 1): f(k, "=KDV+1")})
        y = GeriAlYigini()
        y.calistir(AdDegis(k, ad_anahtari(None, "KDV"), None))
        assert s.sonuc(1, 1) == D.AD
        y.geri_al()
        assert s.sonuc(1, 1) == 3.0

    def test_yeniden_adlandir_formulleri_duzeltir(self):
        k, s, m = kur()
        s2 = k.sayfa_ekle("S2")
        ekle(k, "Fiyat", "=[1]!$A$1:$A$2")
        ekle(k, "Toplamlar", "=TOPLA(Fiyat)")
        s.toplu_yaz({(1, 1): Hucre(1.0), (2, 1): Hucre(2.0), (3, 1): f(k, "=TOPLA(Fiyat)")})
        s2.toplu_yaz({(1, 1): f(k, "=MAK(fiyat)")})
        y = GeriAlYigini()
        y.calistir(AdDegis(k, ad_anahtari(None, "Fiyat"), TanimliAd("Liste", "=[1]!$A$1:$A$2")))
        assert s.hucre(3, 1).formul == "=SUM(Liste)"
        assert s2.hucre(1, 1).formul == "=MAX(Liste)"
        assert k.ad_bul("Toplamlar").formul == "=TOPLA(Liste)" or k.ad_bul("Toplamlar").formul == "=SUM(Liste)"
        assert s.sonuc(3, 1) == 3.0 and s2.sonuc(1, 1) == 2.0
        y.geri_al()
        assert s.hucre(3, 1).formul == "=SUM(Fiyat)" and s2.hucre(1, 1).formul == "=MAX(fiyat)"
        assert k.ad_bul("Fiyat") is not None and k.ad_bul("Liste") is None
        assert s.sonuc(3, 1) == 3.0
        y.yinele()
        assert s.hucre(3, 1).formul == "=SUM(Liste)"

    def test_yeniden_adlandir_golgelenen_sayfaya_dokunmaz(self):
        k, s, m = kur()
        s2 = k.sayfa_ekle("S2")
        ekle(k, "Oran", "=1")
        ekle(k, "Oran", "=2", kapsam=s2.kimlik)           # S2'de gölge
        s.toplu_yaz({(1, 1): f(k, "=Oran")})
        s2.toplu_yaz({(1, 1): f(k, "=Oran")})
        GeriAlYigini().calistir(AdDegis(k, ad_anahtari(None, "Oran"), TanimliAd("Yeni", "=1")))
        assert s.hucre(1, 1).formul == "=Yeni"
        assert s2.hucre(1, 1).formul == "=Oran"

    def test_yeniden_adlandir_yalniz_buyuk_kucuk_formul_degismez(self):
        k, s, m = kur()
        ekle(k, "kdv", "=2")
        s.toplu_yaz({(1, 1): f(k, "=kdv")})
        GeriAlYigini().calistir(AdDegis(k, ad_anahtari(None, "kdv"), TanimliAd("KDV", "=2")))
        assert s.hucre(1, 1).formul == "=kdv"
        assert s.sonuc(1, 1) == 2.0

    def test_adi_degistir_metin_ve_baska_sozcuk(self):
        assert adi_degistir('=A+"A"&Ab+A1', "A", "X") == '=X+"A"&Ab+A1'
        assert adi_degistir("=SUM(B1)", "B", "X") is None
        assert adi_degistir("=Işık*2", "ışık", "Yeni") == "=Yeni*2"

    def test_yinelenen_ad_komutu_kurulamaz(self):
        k, s, m = kur()
        ekle(k, "A", "=1")
        with pytest.raises(AdHatasi):
            AdDegis(k, None, TanimliAd("a", "=2"))


class TestEkleSilTasi:
    def test_satir_ekle_ad_formulunu_duzeltir(self):
        k, s, m = kur()
        ekle(k, "Fiyat", "=[1]!$A$2:$A$3")
        s.toplu_yaz({(2, 1): Hucre(1.0), (3, 1): Hucre(2.0), (5, 1): f(k, "=TOPLA(Fiyat)")})
        y = GeriAlYigini()
        y.calistir(SatirEkleSil(s, 1, 1, True))             # 1. satırın üstüne bir satır
        assert k.ad_bul("Fiyat").formul == "=[1]!$A$3:$A$4"
        assert s.sonuc(6, 1) == 3.0
        y.geri_al()
        assert k.ad_bul("Fiyat").formul == "=[1]!$A$2:$A$3"
        assert s.sonuc(5, 1) == 3.0

    def test_satir_sil_aralik_kuculur_ve_tamamen_silinen_ref(self):
        k, s, m = kur()
        ekle(k, "Bir", "=[1]!$A$2")
        ekle(k, "Iki", "=[1]!$A$2:$A$4")
        y = GeriAlYigini()
        y.calistir(SatirEkleSil(s, 2, 1, False))
        assert k.ad_bul("Bir").formul == "=#REF!"
        assert k.ad_bul("Iki").formul == "=[1]!$A$2:$A$3"
        y.geri_al()
        assert k.ad_bul("Bir").formul == "=[1]!$A$2"
        assert k.ad_bul("Iki").formul == "=[1]!$A$2:$A$4"

    def test_kes_yapistir_adi_kaydirir(self):
        k, s, m = kur()
        ekle(k, "Fiyat", "=[1]!$A$1:$A$2")
        s.toplu_yaz({(1, 1): Hucre(1.0), (2, 1): Hucre(2.0), (4, 3): f(k, "=TOPLA(Fiyat)")})
        y = GeriAlYigini()
        y.calistir(HucreTasi(s, (1, 1, 2, 1), s, (1, 5)))   # A1:A2 → E1:E2
        assert k.ad_bul("Fiyat").formul == "=[1]!$E$1:$E$2"
        assert s.sonuc(4, 3) == 3.0
        y.geri_al()
        assert k.ad_bul("Fiyat").formul == "=[1]!$A$1:$A$2"
        assert s.sonuc(4, 3) == 3.0

    def test_kes_yapistir_kismen_tasinan_ad_degismez(self):
        k, s, m = kur()
        ekle(k, "Fiyat", "=[1]!$A$1:$A$4")
        GeriAlYigini().calistir(HucreTasi(s, (1, 1, 2, 1), s, (1, 5)))
        assert k.ad_bul("Fiyat").formul == "=[1]!$A$1:$A$4"

    def test_kopyala_yapistir_ad_etkilemez(self):
        # kopyala-yapıştır adlara dokunmaz (yalnız kes-yapıştır taşır) — komut yazmadan doğrulama:
        k, s, m = kur()
        ekle(k, "Fiyat", "=[1]!$A$1")
        s.toplu_yaz({(1, 1): Hucre(3.0), (1, 2): Hucre(3.0)})
        assert k.ad_bul("Fiyat").formul == "=[1]!$A$1"


class TestPerformans:
    def test_ad_eklemek_50000_formullu_kitapta_hizli(self):
        import time
        k, s, m = kur()
        yeni = {(r, 1): Hucre(float(r)) for r in range(1, 50001)}
        yeni.update({(r, 2): Hucre(None, "=A%d*2" % r) for r in range(1, 50001)})
        s.toplu_yaz(yeni)
        t = time.perf_counter()
        ekle(k, "Yeni", "=1")
        assert time.perf_counter() - t < 6.0       # ölçüt ≤ 3 sn (makine yükü payı)


class TestXlsx:
    def test_gidis_donus(self, tmp_path):
        from tabcel.io import xlsx_io
        k, s, m = kur()
        s2 = k.sayfa_ekle("Veri Sayfası")
        s.toplu_yaz({(1, 1): Hucre(1.0), (2, 1): Hucre(2.0), (3, 1): f(k, "=TOPLA(Fiyat)+KDV"),
                     (4, 1): f(k, "=Yerel")})
        ekle(k, "Fiyat", "=[1]!$A$1:$A$2")
        ekle(k, "KDV", "=0.18")
        ekle(k, "Yerel", f"=[{s2.kimlik}]!$B$2", kapsam=s2.kimlik)
        ekle(k, "Yerel", "=[1]!$A$1", kapsam=s.kimlik)
        k.ad_yaz(TanimliAd("Gizli", "=1", None, "açıklama", True))
        yol = str(tmp_path / "a.xlsx")
        xlsx_io.kaydet(k, yol)
        k2, rapor = xlsx_io.ac(yol)
        assert not rapor.veri_kaybi
        assert set(k2.adlar) == set(k.adlar)
        t = k2.ad_bul("Gizli")
        assert t.gizli and t.aciklama == "açıklama"
        s_1 = k2.sayfalar[0]
        assert k2.ad_bul("Yerel", s_1.kimlik).formul == "=[1]!$A$1"
        assert k2.ad_bul("Yerel", k2.sayfalar[1].kimlik).kapsam == k2.sayfalar[1].kimlik
        assert s_1.hucre(3, 1).formul == "=SUM(Fiyat)+KDV"
        Hesaplayici(k2)
        assert s_1.sonuc(3, 1) == pytest.approx(3.18)

    def test_dosya_icinde_ingilizce_metin(self, tmp_path):
        from openpyxl import load_workbook
        from tabcel.io import xlsx_io
        k, s, m = kur()
        k.sayfa_ekle("Veri Sayfası")
        ekle(k, "Fiyat", "=[2]!$A$1:$A$2")
        ekle(k, "Oran", "=0.18")
        yol = str(tmp_path / "b.xlsx")
        xlsx_io.kaydet(k, yol)
        wb = load_workbook(yol)
        assert wb.defined_names["Fiyat"].attr_text == "'Veri Sayfası'!$A$1:$A$2"
        assert wb.defined_names["Oran"].attr_text == "0.18"

    def test_cevrilemeyen_ad_atlanir_ve_raporlanir(self, tmp_path):
        from openpyxl import Workbook
        from openpyxl.workbook.defined_name import DefinedName
        from tabcel.io import xlsx_io
        wb = Workbook()
        wb.active.title = "S1"
        wb.defined_names["Iyi"] = DefinedName("Iyi", attr_text="S1!$A$1")
        wb.defined_names["Dis"] = DefinedName("Dis", attr_text="[1]Dis!$A$1")
        wb.defined_names["R1C1"] = DefinedName("R1C1", attr_text="S1!$A$1")
        yol = str(tmp_path / "c.xlsx")
        wb.save(yol)
        k, rapor = xlsx_io.ac(yol)
        assert k.ad_bul("Iyi") is not None and k.ad_bul("Dis") is None and k.ad_bul("R1C1") is None
        assert any("Dis" in u for u in rapor.uyarilar)

    def test_yazdirma_alani_ad_sayilmaz(self, tmp_path):
        from openpyxl import Workbook
        from tabcel.io import xlsx_io
        wb = Workbook()
        wb.active.print_area = "A1:C5"
        yol = str(tmp_path / "d.xlsx")
        wb.save(yol)
        k, rapor = xlsx_io.ac(yol)
        assert k.adlar == {}
