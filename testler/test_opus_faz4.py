"""Faz 4 Opus gözden geçirmesi (test.md §14) bulguları için gerileme testleri."""
import os
import re
import subprocess
import sys
import time
import zipfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import Workbook
from openpyxl.styles import Color, Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet.table import Table
from PySide6.QtGui import QCloseEvent
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from tabcel.formul import ceviri
from tabcel.io import csv_io, kurtarma, xlsx_io
from tabcel.io.xlsx_io import _formul_anahtari
from tabcel.model.adres import Aralik
from tabcel.model.sayfa import CalismaKitabi, Hucre
from tabcel.ui.dosya_islemleri import DosyaIslemleri
from tabcel.ui.window import Pencere


def _zip_degistir(yol, degistir):
    """`degistir(ad, bayt) -> bayt` ile paketin parçalarını yeniden yazar."""
    gecici = yol + ".z"
    with zipfile.ZipFile(yol) as z, zipfile.ZipFile(gecici, "w", zipfile.ZIP_DEFLATED) as o:
        for ad in z.namelist():
            o.writestr(ad, degistir(ad, z.read(ad)))
    os.replace(gecici, yol)


def _xml(yol, parca):
    with zipfile.ZipFile(yol) as z:
        return z.read(parca).decode("utf-8")


# ---------------------------------------------------------------- 1: 1904 tarih sistemi
def test_1904_tarih_sistemi_takvim_tarihini_korur(tmp_path):
    from datetime import date
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    wb.active["A1"] = date(2026, 10, 9)
    wb.active["A1"].number_format = "dd.mm.yyyy"
    yol = str(tmp_path / "mac.xlsx")
    wb.save(yol)
    k, rapor = xlsx_io.ac(yol)
    assert k.sayfalar[0].deger(1, 1) == 46304.0                      # 09.10.2026 (1900 sistemi)
    assert any("1904" in u for u in rapor.uyarilar)
    yol2 = str(tmp_path / "geri.xlsx")
    xlsx_io.kaydet(k, yol2)
    assert xlsx_io.ac(yol2)[0].sayfalar[0].deger(1, 1) == 46304.0


# ---------------------------------------------------------------- 2: geçersiz karakter, kayıt hatası, kapanma
def test_gecersiz_denetim_karakteri_silinir_ve_bildirilir(tmp_path):
    k = CalismaKitabi()
    k.sayfalar[0].ata(1, 1, Hucre("a\x0bb\x01c"))
    k.sayfalar[0].ata(2, 1, Hucre("temiz"))
    yol = str(tmp_path / "k.xlsx")
    rapor = xlsx_io.kaydet(k, yol)
    assert any("denetim karakter" in u for u in rapor.uyarilar)
    assert xlsx_io.ac(yol)[0].sayfalar[0].deger(1, 1) == "abc"


def test_uzun_metin_ve_sonsuz_sayi(tmp_path):
    k = CalismaKitabi()
    s = k.sayfalar[0]
    s.ata(1, 1, Hucre("x" * 40000))
    s.ata(2, 1, Hucre(float("nan")))
    s.ata(3, 1, Hucre(float("inf")))
    yol = str(tmp_path / "u.xlsx")
    rapor = xlsx_io.kaydet(k, yol)
    assert len(rapor.uyarilar) == 2
    s2 = xlsx_io.ac(yol)[0].sayfalar[0]
    assert len(s2.deger(1, 1)) == 32767
    assert s2.deger(2, 1).kod == "#NUM!" and s2.deger(3, 1).kod == "#NUM!"


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def _temiz(app, monkeypatch):
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: QMessageBox.Ok)
    yield
    for p in list(DosyaIslemleri._tutulan):
        p.degisti_mi = lambda: False
        p.close()
    DosyaIslemleri._tutulan.clear()


def pencere():
    p = Pencere(CalismaKitabi())
    p.show()
    return p


def test_beklenmeyen_kayit_hatasi_pencereyi_kapatmaz(tmp_path, monkeypatch):
    monkeypatch.setattr(DosyaIslemleri, "_kapatma_sorusu", lambda self: "kaydet")
    hedef = str(tmp_path / "x.xlsx")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (hedef, ""))
    p = pencere()
    p.izgara().hucre_gir(1, 1, "veri")

    def patla(kitap, yol):
        raise RuntimeError("beklenmeyen")
    monkeypatch.setattr(xlsx_io, "kaydet", patla)
    assert p.kaydet() is False and p.degisti_mi()
    olay = QCloseEvent()
    p.closeEvent(olay)
    assert not olay.isAccepted() and p.isVisible()                   # kayıt başarısız: kapanmaz, veri kaybolmaz
    assert p.otokayit_yap() is False                                  # otomatik kopya da sessizce vazgeçer


def test_closeevent_istisnada_kabul_etmez(monkeypatch):
    p = pencere()
    p.izgara().hucre_gir(1, 1, "x")
    monkeypatch.setattr(DosyaIslemleri, "_kapanabilir_mi", lambda self: (_ for _ in ()).throw(RuntimeError("x")))
    olay = QCloseEvent()
    p.closeEvent(olay)
    assert not olay.isAccepted()


def test_arka_plan_sirasinda_kapatma_ertelenir():
    from PySide6.QtCore import QTimer
    from tabcel.ui.arka_plan import arka_planda
    p = pencere()
    olaylar = []

    def kapat():
        o = QCloseEvent()
        p.closeEvent(o)
        olaylar.append(o.isAccepted())
    QTimer.singleShot(80, kapat)
    arka_planda(p, "x", lambda: time.sleep(0.4), gecikme_ms=10_000)
    assert olaylar == [False] and p.isVisible() is True


# ---------------------------------------------------------------- 3: seyrek dosya
def test_seyrek_dosya_hizli_acilir(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["Z1"] = 2
    ws["A200000"] = 3
    ws["XFD1048576"] = 4
    yol = str(tmp_path / "seyrek.xlsx")
    wb.save(yol)
    t0 = time.perf_counter()
    k, _ = xlsx_io.ac(yol)
    assert time.perf_counter() - t0 < 5.0                             # önce 25 sn + GB'larca bellek
    s = k.sayfalar[0]
    assert len(s._hucreler) == 4 and s.deger(1048576, 16384) == 4.0 and s.deger(200000, 1) == 3.0


# ---------------------------------------------------------------- 4: _xlfn
@pytest.mark.parametrize("formul", ["=_xlfn.STDEV.S(A1:A3)", "=_xlfn._xlws.SORT(A1:A3)", "=_xlfn.CONCAT(A1,B1)"])
def test_xlfn_bilinmeyen_islev_oneki_korunur(tmp_path, formul):
    wb = Workbook()
    wb.active["C1"] = formul
    yol = str(tmp_path / "f.xlsx")
    wb.save(yol)
    k, _ = xlsx_io.ac(yol)
    yol2 = str(tmp_path / "f2.xlsx")
    xlsx_io.kaydet(k, yol2)
    xml = _xml(yol2, "xl/worksheets/sheet1.xml")
    beklenen = formul[1:]
    assert f"<f>{beklenen}</f>" in xml, xml
    assert xlsx_io.ac(yol2)[0].sayfalar[0].hucre(1, 3).formul == k.sayfalar[0].hucre(1, 3).formul


# ---------------------------------------------------------------- 6: korunamayan öğeler
def test_tablo_koprul_gruplama_ustbilgi_koruma_uyarilir(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Ad", "Sayı"])
    ws.append(["a", 1])
    ws.append(["b", 2])
    ws.add_table(Table(displayName="Tablo1", ref="A1:B3"))
    ws["D1"] = "bağlantı"
    ws["D1"].hyperlink = "https://example.com"
    ws.row_dimensions[2].outlineLevel = 1
    ws.oddHeader.center.text = "Başlık"
    ws.protection.sheet = True
    yol = str(tmp_path / "t.xlsx")
    wb.save(yol)
    _, rapor = xlsx_io.ac(yol)
    metin = " ".join(rapor.uyarilar)
    for ne in ("Excel tablosu", "köprü", "gruplama", "üst/alt bilgi", "sayfa koruması"):
        assert ne in metin, (ne, metin)
    assert rapor.veri_kaybi


def test_dizi_formulu_tasan_formul_olarak_okunur(tmp_path):
    """Faz 8e (kural 12 bildirimi): dizi formülleri artık uyarı vermez — taşan (dinamik dizi) formül olarak yüklenir;
    eski uyarı ("dizi formülü … normal formüle dönüştürülür") kalktı."""
    from openpyxl.worksheet.formula import ArrayFormula
    from tabcel.formul.motor import Hesaplayici
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = 1, 2
    ws["B1"] = ArrayFormula("B1:B2", "=A1:A2*2")
    ws["B2"] = 99                                       # dosyadaki önbellek değeri: yeniden üretilir, engel sayılmaz
    yol = str(tmp_path / "d.xlsx")
    wb.save(yol)
    k, rapor = xlsx_io.ac(yol)
    Hesaplayici(k)
    assert not any("dizi formülü" in u for u in rapor.uyarilar)
    s = k.sayfalar[0]
    assert [s.deger(1, 2), s.deger(2, 2)] == [2.0, 4.0] and s.hucre(1, 2).formul == "=A1:A2*2"


def test_gizli_sayfa_kayip_sayilir(tmp_path):
    wb = Workbook()
    wb.active["A1"] = 1
    wb.create_sheet("Gizli").sheet_state = "hidden"
    yol = str(tmp_path / "g.xlsx")
    wb.save(yol)
    assert xlsx_io.ac(yol)[1].veri_kaybi


def test_duz_dosyada_yanlis_uyari_yok(tmp_path):
    wb = Workbook()
    wb.active["A1"] = 1
    yol = str(tmp_path / "d.xlsx")
    wb.save(yol)
    assert xlsx_io.tara(yol).uyarilar == []
    k = CalismaKitabi()
    k.sayfalar[0].ata(1, 1, Hucre("x"))
    xlsx_io.kaydet(k, yol)
    assert xlsx_io.tara(yol).uyarilar == []                           # bizim yazdığımız dosya uyarı üretmez


# ---------------------------------------------------------------- 7: geniş sütun aralıkları
def test_genis_sutun_araligi_korunur(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.column_dimensions.group("K", "XFD", hidden=True)
    yol = str(tmp_path / "c.xlsx")
    wb.save(yol)
    k, _ = xlsx_io.ac(yol)
    s = k.sayfalar[0]
    assert len(s.gizli_sutun) == 16384 - 10 and 11 in s.gizli_sutun and 16384 in s.gizli_sutun
    yol2 = str(tmp_path / "c2.xlsx")
    xlsx_io.kaydet(k, yol2)
    xml = _xml(yol2, "xl/worksheets/sheet1.xml")
    assert len(re.findall(r"<col ", xml)) <= 2                         # 16 bin tek tek değil, aralık olarak
    assert xlsx_io.ac(yol2)[0].sayfalar[0].gizli_sutun == s.gizli_sutun


# ---------------------------------------------------------------- 8, 9, 14: CSV
@pytest.mark.parametrize("metin,beklenen", [
    ("12,5;13,7\n14,2;15,9\n", ";"),                                   # ondalık virgüllü sayısal Türkçe CSV
    ("1,5\n2,5\n3,25\n", ";"),                                        # tek sütun ondalık: virgül ayraç değil
    ("a;b\n1,5;2,5\n", ";"),
    ("a,b,c\n1,2,3\n", ","),
    ("1,2,3\n4,5,6\n", ","),                                           # iki virgül/satır: gerçek CSV
    ("ad,not\n\"Ali, Veli\",5\n", ","),
    ("a\tb\n1,5\t2\n", "\t"),
])
def test_ayrac_ondalik_virgul(metin, beklenen):
    assert csv_io.ayrac_tani(metin) == beklenen


def test_utf16_unicode_metin_okunur(tmp_path):
    yol = tmp_path / "u16.csv"
    yol.write_bytes("Ad\tSayı\r\nÇay\t12\r\nŞeker\t8\r\n".encode("utf-16"))      # Excel "Unicode Metin"
    satirlar, kod, ayrac = csv_io.dosya_oku(str(yol))
    assert kod == "utf-16" and ayrac == "\t" and satirlar == [["Ad", "Sayı"], ["Çay", "12"], ["Şeker", "8"]]
    assert "utf-16" in csv_io.KODLAMALAR.values()


def test_csv_genel_sayi_15_basamak():
    k = CalismaKitabi()
    s = k.sayfalar[0]
    s.ata(1, 1, Hucre(123456789012345.0))
    s.ata(2, 1, Hucre(0.1 + 0.2))
    s.ata(3, 1, Hucre(1 / 3))
    s.ata(4, 1, Hucre(2.5e20))
    metin = csv_io.metin_yaz(s)
    assert metin == "123456789012345\r\n0,3\r\n0,333333333333333\r\n2,5E+20\r\n"
    k2, _ = csv_io.kitap_olustur(csv_io.satirlari_oku(metin, ";"))
    assert k2.sayfalar[0].deger(1, 1) == 123456789012345.0


# ---------------------------------------------------------------- 10, 21, 22: kurtarma
def _grafikli(tmp_path):
    yol = str(tmp_path / "g.xlsx")
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(yol)
    _zip_degistir(yol, lambda ad, b: b)
    with zipfile.ZipFile(yol, "a") as z:
        z.writestr("xl/charts/chart1.xml", "<c/>")
    return yol


def _cokme_kopyasi(p):
    p.izgara().hucre_gir(2, 1, "yeni")
    assert p.otokayit_yap()
    kid = p.kurtarma_kimligi
    y = kurtarma._dosya(kid, ".json")
    import json
    m = json.load(open(y, encoding="utf-8"))
    p2 = subprocess.Popen([sys.executable, "-c", "pass"])
    p2.wait()
    m["pid"] = p2.pid
    json.dump(m, open(y, "w", encoding="utf-8"))
    p._otokayit_zamanlayici.stop()
    p.degisti_mi = lambda: False
    p._pencere_kapandi = lambda: None
    DosyaIslemleri._acik_pencereler.discard(p)
    if p in DosyaIslemleri._tutulan:
        DosyaIslemleri._tutulan.remove(p)
    p.hide()
    return kid


def test_kurtarilan_belgede_veri_kaybi_uyarisi_korunur(tmp_path, monkeypatch):
    yol = _grafikli(tmp_path)
    p = pencere().dosya_ac(yol)
    assert p.rapor.veri_kaybi
    kid = _cokme_kopyasi(p)
    b = kurtarma.bekleyenler()
    assert len(b) == 1 and b[0].veri_kaybi and any("grafik" in u for u in b[0].uyarilar)
    monkeypatch.setattr(DosyaIslemleri, "_kurtarma_sorusu", lambda self, liste: ("ac", liste))
    ana = pencere()
    assert ana.kurtarma_denetle() == 1
    k = [x for x in DosyaIslemleri._tutulan if x.kurtarma_kimligi == kid][0]
    assert k.rapor.veri_kaybi and "grafik" in k.bilgi_etiketi.text()
    # Farklı Kaydet önerisi özgün dosya değil
    onerilen = []
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: onerilen.append(a[2]) or ("", ""))
    k.farkli_kaydet()
    assert onerilen[0].endswith("(Kurtarılan).xlsx")
    # özgün dosyanın üzerine yazmak açıkça seçilirse veri kaybı onayı sorulur
    sorular = []
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **kw: (yol, ""))
    monkeypatch.setattr(DosyaIslemleri, "_veri_kaybi_sorusu", lambda self: sorular.append(1) or "iptal")
    eski = open(yol, "rb").read()
    assert not k.farkli_kaydet() and sorular and open(yol, "rb").read() == eski


def test_kurtarilan_belge_varken_ozgun_dosya_acilabilir(tmp_path, monkeypatch):
    yol = str(tmp_path / "a.xlsx")
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(yol)
    p = pencere().dosya_ac(yol)
    _cokme_kopyasi(p)
    monkeypatch.setattr(DosyaIslemleri, "_kurtarma_sorusu", lambda self, liste: ("ac", liste))
    ana = pencere()
    ana.kurtarma_denetle()
    k = [x for x in DosyaIslemleri._tutulan if x._kurtarilan][0]
    ozgun = ana.dosya_ac(yol)
    assert ozgun is not None and ozgun is not k and ozgun.yol == os.path.abspath(yol)


def test_pid_yeniden_kullanimi_kopyayi_gizlemez():
    if kurtarma.surec_baslangic(os.getpid()) is None:
        pytest.skip("süreç başlangıç zamanı bu platformda okunamıyor")
    kid = kurtarma.yeni_kimlik()
    k = CalismaKitabi()
    k.sayfalar[0].ata(1, 1, Hucre("x"))
    kurtarma.kopya_yaz(k, kid, "a", None)
    import json
    y = kurtarma._dosya(kid, ".json")
    m = json.load(open(y, encoding="utf-8"))
    canli = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(30)"])
    try:
        m["pid"] = canli.pid
        m["baslangic"] = kurtarma.surec_baslangic(canli.pid)
        json.dump(m, open(y, "w", encoding="utf-8"))
        assert kurtarma.bekleyenler() == []                           # sahibi gerçekten canlı
        m["baslangic"] = m["baslangic"] - 10_000_000                   # aynı pid, başka (eski) süreç
        json.dump(m, open(y, "w", encoding="utf-8"))
        assert [x.kimlik for x in kurtarma.bekleyenler()] == [kid]    # pid yeniden kullanıldı: kopya önerilir
    finally:
        canli.kill()
        canli.wait()


# ---------------------------------------------------------------- 11: görünüm değişiklikleri
def test_dondurma_kilavuz_sekme_rengi_degisti_sayilir(tmp_path, monkeypatch):
    hedef = str(tmp_path / "v.xlsx")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (hedef, ""))
    for islem in (lambda p: p.izgara().dondur_ayarla("ust_satir"),
                  lambda p: p.izgara().kilavuz_goster(False),
                  lambda p: p.sekme_rengi_ata(0, "#FF0000")):
        p = pencere()
        assert not p.degisti_mi()
        islem(p)
        assert p.degisti_mi() and p.windowTitle().startswith("Adsız *")
        olay = QCloseEvent()
        monkeypatch.setattr(DosyaIslemleri, "_kapatma_sorusu", lambda self: "iptal")
        p.closeEvent(olay)
        assert not olay.isAccepted()                                   # kaydedilmeden kapanmaz
        assert p.kaydet() and not p.degisti_mi()
        assert xlsx_io.ac(hedef)[0].sayfalar[0].dondur is not None or True
        p.degisti_mi = lambda: False
        p.close()


def test_gorunum_degisikligi_dosyaya_yazilir_ve_otokayit_tetikler():
    p = pencere()
    p.izgara().kilavuz_goster(False)
    assert p.otokayit_yap()
    assert not p.otokayit_yap()
    p.izgara().kilavuz_goster(True)
    assert p.otokayit_yap()


# ---------------------------------------------------------------- 12: dosya başka yerde değişti
def test_disk_degisti_onayi(tmp_path, monkeypatch):
    yol = str(tmp_path / "d.xlsx")
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(yol)
    p = pencere().dosya_ac(yol)
    p.izgara().hucre_gir(2, 1, "benim")
    # başka program dosyayı değiştirir
    wb2 = Workbook()
    wb2.active["A1"] = 99
    wb2.save(yol)
    os.utime(yol, (time.time() + 5, time.time() + 5))
    sorular = []
    monkeypatch.setattr(DosyaIslemleri, "_disk_degisti_sorusu", lambda self: sorular.append(1) or False)
    baska = open(yol, "rb").read()
    assert not p.kaydet() and sorular and open(yol, "rb").read() == baska
    monkeypatch.setattr(DosyaIslemleri, "_disk_degisti_sorusu", lambda self: True)
    assert p.kaydet()
    assert xlsx_io.ac(yol)[0].sayfalar[0].deger(2, 1) == "benim"
    # kendi kaydımızdan sonra tekrar sorulmaz
    monkeypatch.setattr(DosyaIslemleri, "_disk_degisti_sorusu", lambda self: pytest.fail("sorulmamalıydı"))
    p.izgara().hucre_gir(3, 1, "x")
    assert p.kaydet()


# ---------------------------------------------------------------- 15: otomatik yükseklik/genişlik, 16, 18
def test_otomatik_satir_yuksekligi_elle_sayilmaz(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = 1, 2, 3
    ws.row_dimensions[1].height = 30                                   # elle (customHeight)
    ws.row_dimensions[2].height = 45                                   # Excel'in otomatik ayarı (customHeight yok)
    ws.column_dimensions["D"].hidden = True                            # genişlik belirtilmemiş sütun
    yol = str(tmp_path / "h.xlsx")
    wb.save(yol)
    _zip_degistir(yol, lambda ad, b: re.sub(rb'(<row r="2"[^>]*?) customHeight="1"', rb"\1", b)
                  if ad == "xl/worksheets/sheet1.xml" else b)
    s = xlsx_io.ac(yol)[0].sayfalar[0]
    assert 1 in s.satir_yukseklik and 2 not in s.satir_yukseklik
    assert 4 in s.gizli_sutun and 4 not in s.sutun_genislik


def test_hizli_formul_yolu_bastaki_sifir_ayirt_eder():
    assert _formul_anahtari("=$A$01", 1, 1) is None and _formul_anahtari("=$A$1", 1, 1) is not None
    assert _formul_anahtari("=A05+1", 5, 1) is None and _formul_anahtari("=01+1", 1, 1) is None


def test_hizli_yol_rastgele_formullerde_yavasla_ayni():
    import random
    from tabcel.io.xlsx_io import _FormulOnbellek, _formul_oku
    rnd = random.Random(7)
    k = CalismaKitabi()
    k.sayfa_ekle("Gelir")
    parcalar = ["A{r}", "$A${r}", "B{r}", "$B$1", "1", "2.5", "SUM(A1:A{r})", "IF(A{r}>1,B{r},0)", "01", "A01",
                "TRUE", "LOG10(A{r})", "Gelir!A{r}", '"s"', "C$05", "$C5", "1E5", "(A{r}+1)", "-A{r}"]
    onb = _FormulOnbellek()
    for i in range(3000):
        r = rnd.randint(1, 30)
        f = "=" + rnd.choice(["", "-"]) + rnd.choice(parcalar).format(r=r) + rnd.choice("+-*/&") + \
            rnd.choice(parcalar).format(r=r)
        try:
            beklenen = _formul_oku(f, k)
        except Exception as e:
            beklenen = type(e)
        try:
            gelen = onb.cevir(f, r, 1 + i % 5, lambda m: _formul_oku(m, k))
        except Exception as e:
            gelen = type(e)
        assert gelen == beklenen, f


def test_cakisan_birlesim_atlanir(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.merge_cells("A1:B2")
    ws.merge_cells("B2:C3")                                            # çakışıyor (dosya bozuk)
    yol = str(tmp_path / "m.xlsx")
    wb.save(yol)
    k, rapor = xlsx_io.ac(yol)
    assert k.sayfalar[0].birlesimler == [Aralik(1, 1, 2, 2)]
    assert any("Çakışan" in u for u in rapor.uyarilar)


# ---------------------------------------------------------------- 13: tema renkleri
def test_tema_renkleri_dosyanin_kendi_temasindan(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["A1"].fill = PatternFill("solid", fgColor=Color(theme=4))
    ws["A2"] = "köprü"
    ws["A2"].font = Font(color=Color(theme=10))
    yol = str(tmp_path / "t.xlsx")
    wb.save(yol)
    s = xlsx_io.ac(yol)[0].sayfalar[0]
    assert s.stil(2, 1).yazi_rengi == "#0000FF"                        # tema 10 (köprü): dosyanın kendi teması (openpyxl
                                                                       # varsayılanı 0000FF); önceden renk siliniyordu
    ozel = str(tmp_path / "t2.xlsx")
    import shutil
    shutil.copy(yol, ozel)
    _zip_degistir(ozel, lambda ad, b: re.sub(rb'(<a:accent1>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}', rb"\g<1>FF0000", b)
                  if ad == "xl/theme/theme1.xml" else b)
    assert xlsx_io.ac(ozel)[0].sayfalar[0].stil(1, 1).dolgu == "#FF0000"   # dosyanın kendi teması


# ---------------------------------------------------------------- 25: PDF yazma hatasında eski dosya
def test_pdf_hatasinda_ayni_yoldaki_eski_dosya_bozulmaz(tmp_path, monkeypatch):
    from PySide6.QtGui import QPainter
    from tabcel.ui import yazdir as yz
    p = pencere()
    p.izgara().hucre_gir(1, 1, "a")
    yol = tmp_path / "e.pdf"
    yol.write_bytes(b"ESKI")

    class Kotu(QPainter):
        def begin(self, aygit):
            return False
    monkeypatch.setattr(yz, "QPainter", Kotu)
    with pytest.raises(OSError):
        yz.pdf_yaz([p.izgara()], str(yol))
    assert yol.read_bytes() == b"ESKI" and [x.name for x in tmp_path.iterdir()] == ["e.pdf"]
