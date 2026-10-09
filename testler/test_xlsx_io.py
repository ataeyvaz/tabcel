"""Faz 4 / 4a: .xlsx gidiş-dönüş (aç → kaydet → aç): değer, formül, biçim, boyut, birleşim, dondurma, çok sayfa."""
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from tabcel.formul import ceviri
from tabcel.io import xlsx_io
from tabcel.model.adres import Aralik
from tabcel.model.deger import BOL0
from tabcel.model.komutlar import bicim_komutu
from tabcel.model.sayfa import CalismaKitabi, Hucre
from tabcel.model.stil import StilYama


def _kitap():
    k = CalismaKitabi()
    return k, k.sayfalar[0]


def _yaz(s, ar, **yama):
    bicim_komutu(s, [ar], StilYama(**yama)).yap()


def _gidis_donus(k, tmp_path, ad="a.xlsx"):
    yol = str(tmp_path / ad)
    xlsx_io.kaydet(k, yol)
    return xlsx_io.ac(yol)


def test_degerler_ve_formuller(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre("Merhaba"))
    s.ata(1, 2, Hucre(3.5))
    s.ata(1, 3, Hucre(True))
    s.ata(1, 4, Hucre("=metin"))            # '=' ile başlayan DÜZ metin formül olmamalı
    s.ata(2, 1, Hucre(None, ceviri.iceri("=SUM(B1:B1)*2", k)))
    s.ata(2, 2, Hucre(BOL0))
    k2, rapor = _gidis_donus(k, tmp_path)
    s2 = k2.sayfalar[0]
    assert s2.hucre(1, 1).deger == "Merhaba" and s2.hucre(1, 2).deger == 3.5
    assert s2.hucre(1, 3).deger is True
    assert s2.hucre(1, 4).deger == "=metin" and s2.hucre(1, 4).formul is None
    assert s2.hucre(2, 1).formul == "=SUM(B1:B1)*2"
    assert s2.hucre(2, 2).deger == BOL0
    assert not rapor.veri_kaybi


def test_sayi_bicimi_ve_tarih(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre(46037.0))
    _yaz(s, Aralik(1, 1, 1, 1), sayi_bicimi="dd.mm.yyyy")
    s.ata(2, 1, Hucre(0.25))
    _yaz(s, Aralik(2, 1, 2, 1), sayi_bicimi="0.00%")
    k2, _ = _gidis_donus(k, tmp_path)
    s2 = k2.sayfalar[0]
    assert s2.deger(1, 1) == 46037.0 and s2.stil(1, 1).sayi_bicimi == "dd.mm.yyyy"
    assert s2.deger(2, 1) == 0.25 and s2.stil(2, 1).sayi_bicimi == "0.00%"
    ws = load_workbook(str(tmp_path / "a.xlsx")).active          # dosyada gerçekten tarih olarak yazılmış
    assert ws["A1"].number_format == "dd.mm.yyyy" and ws["A1"].is_date


def test_stil_gidis_donus(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre("x"))
    _yaz(s, Aralik(1, 1, 1, 1), kalin=True, italik=True, ustu_cizili=True, alti_cizili="cift",
         yazi_tipi="Arial", boyut=14.0, yazi_rengi="#FF0000", dolgu="#00FF00",
         yatay="orta", dikey="ust", girinti=2, kaydir=True,
         kenar={"ust": ("kalin", "#112233"), "alt": ("cift", "#000000"), "sol": ("kesikli", "#FF00FF")})
    k2, _ = _gidis_donus(k, tmp_path)
    assert k.sayfalar[0].stil(1, 1) == k2.sayfalar[0].stil(1, 1)


def test_varsayilan_font_none_kalir(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre("x"))
    _yaz(s, Aralik(1, 1, 1, 1), kalin=True)
    k2, _ = _gidis_donus(k, tmp_path)
    st = k2.sayfalar[0].stil(1, 1)
    assert st.yazi_tipi is None and st.boyut is None and st.kalin


def test_olculer_gizli_birlesim_dondurma(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre("Başlık"))
    s.birlesimleri_yaz([Aralik(1, 1, 1, 3)], [])
    s.sutun_genislik[2] = 120
    s.satir_yukseklik[3] = 40
    s.gizli_sutun.add(5)
    s.gizli_satir.add(7)
    s.gorunum_yaz(dondur=(2, 2), kilavuz=False)
    s.sekme_rengi = "#217346"
    k2, _ = _gidis_donus(k, tmp_path)
    s2 = k2.sayfalar[0]
    assert s2.birlesimler == [Aralik(1, 1, 1, 3)]
    assert abs(s2.sutun_genislik[2] - 120) <= 7 and abs(s2.satir_yukseklik[3] - 40) <= 2
    assert 5 in s2.gizli_sutun and 7 in s2.gizli_satir
    assert s2.dondur == (2, 2) and s2.kilavuz is False
    assert s2.sekme_rengi == "#217346"


def test_satir_sutun_stili(tmp_path):
    k, s = _kitap()
    s.ata(5, 2, Hucre("x"))
    _yaz(s, Aralik(1, 2, 1_048_576, 2), dolgu="#FF0000")        # B sütunu
    _yaz(s, Aralik(3, 1, 3, 16384), kalin=True)                  # 3. satır
    k2, _ = _gidis_donus(k, tmp_path)
    s2 = k2.sayfalar[0]
    assert s2.stil(100, 2).dolgu == "#FF0000"
    assert s2.stil(3, 5).kalin


def test_cok_sayfa_ve_sayfalar_arasi_formul(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre(10.0))
    k.yeniden_adlandir(s, "Gelir")
    s2 = k.sayfa_ekle("Özet")
    s2.ata(1, 1, Hucre(None, ceviri.iceri("=Gelir!A1*2", k)))
    k3, _ = _gidis_donus(k, tmp_path)
    assert [x.ad for x in k3.sayfalar] == ["Gelir", "Özet"]
    assert ceviri.disa(k3.sayfalar[1].hucre(1, 1).formul, k3) == "=Gelir!A1*2"
    ws = load_workbook(str(tmp_path / "a.xlsx"))["Özet"]
    assert ws["A1"].value == "=Gelir!A1*2"


def test_xlfn_oneki(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre(None, ceviri.iceri('=IFS(1>2,"a",TRUE,"IFS(")', k)))
    s.ata(2, 1, Hucre(None, ceviri.iceri("=XLOOKUP(1,A1:A3,B1:B3)", k)))
    k2, _ = _gidis_donus(k, tmp_path)
    ws = load_workbook(str(tmp_path / "a.xlsx")).active
    assert ws["A1"].value.startswith("=_xlfn.IFS(") and '"IFS("' in ws["A1"].value
    assert ws["A2"].value.startswith("=_xlfn.XLOOKUP(")
    assert k2.sayfalar[0].hucre(1, 1).formul == k.sayfalar[0].hucre(1, 1).formul
    assert k2.sayfalar[0].hucre(2, 1).formul == k.sayfalar[0].hucre(2, 1).formul


def test_baska_programda_olusturulmus_dosya(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Veri"
    ws["A1"] = "Ad"
    ws["B1"] = 5
    ws["B2"] = "=B1*2"
    ws["C1"] = date(2026, 10, 12)
    ws["C1"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 20
    ws.merge_cells("A3:B3")
    yol = str(tmp_path / "b.xlsx")
    wb.save(yol)
    k, _ = xlsx_io.ac(yol)
    s = k.sayfalar[0]
    assert s.ad == "Veri" and s.deger(1, 2) == 5.0 and s.hucre(2, 2).formul == "=B1*2"
    assert s.deger(1, 3) == 46307.0                              # 12.10.2026
    assert s.stil(1, 3).sayi_bicimi == "yyyy-mm-dd"
    assert s.birlesimler == [Aralik(3, 1, 3, 2)]
    assert abs(s.sutun_genislik[1] - 145) <= 8                   # 20 karakter ≈ 145 px


def test_veri_kaybi_taramasi(tmp_path):
    wb = Workbook()
    wb.active["A1"] = 1
    yol = str(tmp_path / "c.xlsx")
    wb.save(yol)
    yol2 = str(tmp_path / "d.xlsx")                              # grafikli paket taklidi
    with zipfile.ZipFile(yol) as z, zipfile.ZipFile(yol2, "w") as o:
        for ad in z.namelist():
            o.writestr(ad, z.read(ad))
        o.writestr("xl/charts/chart1.xml", "<c/>")
    k, rapor = xlsx_io.ac(yol2)
    assert rapor.veri_kaybi and any("grafik" in u for u in rapor.uyarilar)
    assert k.sayfalar[0].deger(1, 1) == 1.0                      # veri korunur
    assert not xlsx_io.tara(yol).veri_kaybi


def test_kosullu_bicim_ve_dogrulama_uyarisi(tmp_path):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.conditional_formatting.add("A1:A5", CellIsRule(operator="greaterThan", formula=["3"]))
    dv = DataValidation(type="list", formula1='"a,b"')
    ws.add_data_validation(dv)
    dv.add("B1")
    yol = str(tmp_path / "e.xlsx")
    wb.save(yol)
    _, rapor = xlsx_io.ac(yol)
    assert any("koşullu biçim" in u for u in rapor.uyarilar)
    assert any("veri doğrulama" in u for u in rapor.uyarilar)


def test_bozuk_dosya_anlasilir_hata(tmp_path):
    yol = tmp_path / "x.xlsx"
    yol.write_bytes(b"bu bir zip degil")
    with pytest.raises(xlsx_io.XlsxHatasi):
        xlsx_io.ac(str(yol))
    with pytest.raises(xlsx_io.XlsxHatasi):
        xlsx_io.ac(str(tmp_path / "yok.xlsx"))


def test_kayit_gecici_dosya_kalmaz(tmp_path):
    k, s = _kitap()
    s.ata(1, 1, Hucre(1.0))
    yol = str(tmp_path / "f.xlsx")
    xlsx_io.kaydet(k, yol)
    xlsx_io.kaydet(k, yol)                                       # üzerine yazma
    assert [p.name for p in tmp_path.iterdir()] == ["f.xlsx"]


def test_kayit_hatasinda_eski_dosya_bozulmaz(tmp_path, monkeypatch):
    k, s = _kitap()
    s.ata(1, 1, Hucre(1.0))
    yol = str(tmp_path / "g.xlsx")
    xlsx_io.kaydet(k, yol)
    eski = open(yol, "rb").read()

    def patla(self, *a, **kw):
        open(a[0], "wb").write(b"yarim")                         # yarıda kalan yazım
        raise OSError("disk dolu")
    monkeypatch.setattr(Workbook, "save", patla)
    with pytest.raises(xlsx_io.XlsxHatasi):
        xlsx_io.kaydet(k, yol)
    assert open(yol, "rb").read() == eski
    assert [p.name for p in tmp_path.iterdir()] == ["g.xlsx"]


def test_gizli_sayfa_uyarisi(tmp_path):
    wb = Workbook()
    wb.active["A1"] = 1
    gizli = wb.create_sheet("Gizli")
    gizli.sheet_state = "hidden"
    yol = str(tmp_path / "h.xlsx")
    wb.save(yol)
    k, rapor = xlsx_io.ac(yol)
    assert [x.ad for x in k.sayfalar] == ["Sheet", "Gizli"]
    assert any("Gizli" in u for u in rapor.uyarilar)


# ---------------------------------------------------------------- 4f: hızlı formül çevirisi
def test_hizli_formul_yolu_yavasla_ayni_sonuc():
    from tabcel.io.xlsx_io import _FormulOnbellek, _formul_anahtari, _formul_oku
    k, s = _kitap()
    k.sayfa_ekle("Gelir")
    formuller = []
    for r in range(1, 40):
        formuller += [f"=A{r}+B{r}*3", f"=SUM(A1:A{r})", f"=$A$1+C{r}", f"=Gelir!A{r}", f"=sum(a{r}:b{r})",
                      f"=IF(A{r}>1,\"x\",\"y\")", f"=LOG10(A{r})", f"=_xlfn.IFS(A{r}>1,1,TRUE,2)",
                      f"=A{r} + B{r}", f"=-A{r}^2%", f"=TOPLA(A{r})", f"=A{r}&B{r}", f"=1.5*A{r}"]
    onb = _FormulOnbellek()
    for i, f in enumerate(formuller):
        r, c = 1 + i // 13, 1 + i % 13
        try:
            beklenen = _formul_oku(f, k)
        except Exception as e:
            beklenen = type(e)
        try:
            gelen = onb.cevir(f, r, c, lambda m: _formul_oku(m, k))
        except Exception as e:
            gelen = type(e)
        assert gelen == beklenen, (f, gelen, beklenen)
    assert _formul_anahtari("=A5+B5", 5, 1) == _formul_anahtari("=A6+B6", 6, 1)       # bağıl yapı aynı
    assert _formul_anahtari("=A5+B5", 5, 1) != _formul_anahtari("=A6+B5", 6, 1)
    assert _formul_anahtari("=$A$1+B5", 5, 1) != _formul_anahtari("=$A$2+B6", 6, 1)   # mutlak başvuru farklı
    assert _formul_anahtari("=Gelir!A1", 1, 1) is None and _formul_anahtari('=IF(A1,"x")', 1, 1) is None


def test_buyuk_dosya_gidis_donus_ayni(tmp_path):
    k, s = _kitap()
    yeni = {}
    for r in range(1, 301):
        yeni[(r, 1)] = Hucre(float(r))
        yeni[(r, 2)] = Hucre(None, ceviri.iceri(f"=A{r}*2+IF(A{r}>10,1,0)", k))
        yeni[(r, 3)] = Hucre(None, ceviri.iceri(f"=SUM($A$1:A{r})", k))
    s.toplu_yaz(yeni)
    k2, _ = _gidis_donus(k, tmp_path)
    for anahtar, h in yeni.items():
        assert k2.sayfalar[0].hucre(*anahtar) == h
