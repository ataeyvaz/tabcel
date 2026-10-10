"""Faz 5 / 5h: koşullu biçim arayüzü (çizim, menü, pencereler, yönetici, geçersiz kılma) + .xlsx — ekransız."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtGui import QColor
from PySide6.QtWidgets import QApplication

from tabcel.io import xlsx_io
from tabcel.model.adres import Aralik
from tabcel.model.kosullu import KosulluKural, RENK_OLCEKLERI
from tabcel.model.sayfa import CalismaKitabi, Hucre
from tabcel.model.stil import StilYama
from tabcel.ui.grid import Izgara
from tabcel.ui.kosullu_pencereleri import (
    HAZIR_BICIMLER, KuralPenceresi, KuralYoneticisi, OzelBicimPenceresi, deger_formulu, deger_gosterimi,
    kural_aciklamasi,
)
from tabcel.ui.window import Pencere

A = (Aralik(1, 1, 8, 1),)
KIRMIZI = StilYama(dolgu="#FF0000")


@pytest.fixture(scope="session")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def pencere(app):
    Izgara._pano_durumu = None
    k = CalismaKitabi()
    s = k.sayfalar[0]
    for r in range(1, 9):
        s.ata(r, 1, Hucre(float(r * 10)))
    p = Pencere(k)
    p.resize(900, 500)
    p.show()
    app.processEvents()
    yield p
    p.hide()


def piksel(g, r, c, dx=3, dy=3):
    img = g.viewport().grab().toImage()
    ar = g.hucre_dikdortgeni(r, c)
    return QColor(img.pixel(ar.x() + dx, ar.y() + dy))


def kural(**kw):
    return KosulluKural((), **kw)


# ----------------------------------------------------------------------------- çizim
def test_kural_hucreyi_boyar_geri_al_ve_olay_ile_gecersiz(pencere, app):
    g = pencere.izgara()
    g.sec((1, 1), (8, 1))
    assert pencere.kosullu_kural_ac(sonuc=kural(tur="hucre_degeri", islec="buyuk", formul1="=50", bicim=KIRMIZI))
    assert piksel(g, 6, 1).red() > 200 and piksel(g, 6, 1).green() < 60          # 60 > 50 boyandı
    assert piksel(g, 2, 1).green() > 200                                          # 20: boyanmadı (beyaz)
    g.sayfa.ata(2, 1, Hucre(99.0))                                               # değer değişince önbellek geçersiz
    app.processEvents()
    assert piksel(g, 2, 1).red() > 200 and piksel(g, 2, 1).green() < 60
    g.geri_al()                                                                   # son komut kural ekleme değil: değer yazımı yok (doğrudan ata)
    g.geri_al()
    assert piksel(g, 6, 1).green() > 200 and g.sayfa.kosullu == []


def test_renk_olcegi_ve_veri_cubugu_cizilir(pencere, app):
    g = pencere.izgara()
    g.sec((1, 1), (8, 1))
    assert pencere.kosullu_hazir("renk_olcegi", "yesil_beyaz")
    assert piksel(g, 8, 1).red() < 200 and piksel(g, 8, 1).green() > piksel(g, 8, 1).red()    # en büyük: yeşil
    assert piksel(g, 1, 1).red() > 240                                                         # en küçük: beyaza yakın
    g.geri_al()
    assert pencere.kosullu_hazir("veri_cubugu", "duz_mavi")
    assert piksel(g, 8, 1, dx=10).blue() > piksel(g, 8, 1, dx=10).red()                       # satır 8: uzun mavi çubuk
    assert piksel(g, 1, 1, dx=40).red() > 240                                                 # satır 1: çubuk kısa, sağı boş


def test_yalniz_cubuk_metni_gizler(pencere):
    g = pencere.izgara()
    g.sec((1, 1), (8, 1))
    pencere.kosullu_kural_ac(sonuc=KosulluKural((), tur="veri_cubugu", bicim=None, yalniz_cubuk=True))
    metinler, _ = g._yerlesim(list(g._gorunen_sutunlar()), list(g._gorunen_satirlar()), g.fontMetrics(), 4)
    assert not [m for m in metinler if m[0].y() >= g.hucre_dikdortgeni(1, 1).y()
                and m[0].x() == g.hucre_dikdortgeni(1, 1).x()]
    assert g._cubuklar


def test_birlesik_hucre_boyanir(pencere):
    g = pencere.izgara()
    g.sec((2, 1), (3, 1))
    g.birlestir("birlestir", onayli=True)
    g.sec((2, 1))
    pencere.kosullu_kural_ac(sonuc=kural(tur="hucre_degeri", islec="buyuk", formul1="=5", bicim=KIRMIZI))
    c = piksel(g, 2, 1)
    assert c.red() > 200 and c.green() < 60


def test_baska_sayfadaki_degisiklik_gecersiz_kilar(pencere, app):
    k = pencere.kitap
    s2 = k.sayfa_ekle("S2")
    app.processEvents()
    pencere.sekmeler.setCurrentIndex(0)                       # pencere yeni sayfaya geçmişti: ilk sayfaya dön
    app.processEvents()
    g = pencere.izgara()
    g.sec((1, 1))
    pencere.kosullu_kural_ac(sonuc=kural(tur="formul", formul1=f"=[{s2.kimlik}]!$A$1>0", bicim=KIRMIZI))
    assert piksel(g, 1, 1).green() > 200
    s2.ata(1, 1, Hucre(5.0))
    app.processEvents()
    assert piksel(g, 1, 1).green() < 60


# ----------------------------------------------------------------------------- menü ve komutlar
def test_menu_ogeleri(pencere):
    e = pencere.kosullu_eylemleri
    for ad in ("buyuk", "kucuk", "arasinda", "esit", "metin", "yinelenen", "ilk10", "ilk10y", "son10", "son10y",
               "ort_ustu", "ort_alti", "cubuk_mavi", "cubuk_duz_mor", "olcek_yesil_sari_kirmizi", "olcek_beyaz_kirmizi",
               "yeni", "temizle_secim", "temizle_sayfa", "yonet"):
        assert ad in e, ad
    assert len([a for a in e if a.startswith("cubuk_")]) == 12 and len([a for a in e if a.startswith("olcek_")]) == 4


def test_yeni_kural_en_yuksek_oncelik_ve_capa_kaymasi(pencere):
    g = pencere.izgara()
    g.sec((1, 1), (8, 1))
    pencere.kosullu_kural_ac(sonuc=kural(tur="bos", bicim=KIRMIZI))
    pencere.kosullu_kural_ac(sonuc=kural(tur="hata", bicim=KIRMIZI))
    assert [k.tur for k in g.sayfa.kosullu] == ["hata", "bos"]
    g.sec((5, 1), (3, 1))                                                    # etkin A5, çapa A3
    pencere.kosullu_kural_ac(sonuc=kural(tur="formul", formul1="=A5>1", bicim=KIRMIZI))
    assert g.sayfa.kosullu[0].formul1 == "=A3>1" and g.sayfa.kosullu[0].araliklar == (Aralik(3, 1, 5, 1),)
    g.geri_al()
    assert [k.tur for k in g.sayfa.kosullu] == ["hata", "bos"]


def test_temizle_secim_ve_sayfa_ve_tumunu_temizle(pencere):
    g = pencere.izgara()
    g.sec((1, 1), (4, 1))
    pencere.kosullu_kural_ac(sonuc=kural(tur="bos", bicim=KIRMIZI))
    g.sec((6, 1), (8, 1))
    pencere.kosullu_kural_ac(sonuc=kural(tur="hata", bicim=KIRMIZI))
    g.sec((1, 1), (2, 1))
    assert pencere.kosullu_temizle(False)
    assert {a for k in g.sayfa.kosullu for a in k.araliklar} == {Aralik(3, 1, 4, 1), Aralik(6, 1, 8, 1)}
    g.geri_al()
    assert len(g.sayfa.kosullu) == 2
    assert pencere.kosullu_temizle(True) and g.sayfa.kosullu == []
    assert not pencere.kosullu_temizle(True)
    g.geri_al()
    g.sec((1, 1), (8, 1))
    g.tumunu_temizle()                                                        # Excel Tümünü Temizle koşullu biçimi de siler
    assert g.sayfa.kosullu == []
    g.geri_al()
    assert len(g.sayfa.kosullu) == 2


def test_yonet_sonucu_uygular_ve_siralama(pencere):
    g = pencere.izgara()
    g.sec((1, 1), (8, 1))
    a, b = kural(tur="bos", bicim=KIRMIZI), kural(tur="hata", bicim=KIRMIZI)
    pencere.kosullu_kural_ac(sonuc=a)
    pencere.kosullu_kural_ac(sonuc=b)
    eski = list(g.sayfa.kosullu)
    assert pencere.kosullu_yonet(list(reversed(eski)))
    assert g.sayfa.kosullu == list(reversed(eski))
    assert not pencere.kosullu_yonet(list(reversed(eski)))


# ----------------------------------------------------------------------------- pencereler
def test_deger_donusumleri(app):
    k = CalismaKitabi()
    assert deger_formulu("2,5", k) == "=2.5" and deger_formulu("=A1*2", k) == "=A1*2"
    assert deger_formulu('Ali "x"', k) == '="Ali ""x"""' and deger_formulu("1.1.2026", k) == "=46023"
    with pytest.raises(ValueError):
        deger_formulu("  ", k)
    with pytest.raises(ValueError):
        deger_formulu("=TOPLA(", k)
    assert deger_gosterimi("=2.5", k) == "2,5" and deger_gosterimi('="Ali"', k) == "Ali"
    assert deger_gosterimi("=SUM(A1:A3)", k) == "=TOPLA(A1:A3)" and deger_gosterimi(None, k) == ""


def test_kural_penceresi_turleri_ve_geri_yukleme(app):
    k = CalismaKitabi()
    p = KuralPenceresi(k, tur="hucre_degeri", islec="arasinda")
    p.show()
    assert p.deger2.isVisible() and p.islec_kutusu.currentData() == "arasinda"
    p.deger1.setText("2,5")
    p.deger2.setText("=B1")
    p.bicim_kutusu.setCurrentIndex(2)
    kr = p.kural_olustur()
    assert (kr.tur, kr.islec, kr.formul1, kr.formul2, kr.bicim) == ("hucre_degeri", "arasinda", "=2.5", "=B1",
                                                                      HAZIR_BICIMLER[2][1])
    geri = KuralPenceresi(k, kural=kr)
    assert geri.deger1.text() == "2,5" and geri.deger2.text() == "=B1" and geri.bicim_kutusu.currentIndex() == 2
    p.close()
    geri.close()


@pytest.mark.parametrize("tur,ayar,beklenen", [
    ("metin", lambda p: p.metin.setText("ağ"), dict(tur="metin", islec="icerir", metin="ağ")),
    ("bos", lambda p: None, dict(tur="bos")),
    ("hata_degil", lambda p: None, dict(tur="hata_degil")),
    ("benzersiz", lambda p: None, dict(tur="yinelenen", benzersiz=True)),
    ("son", lambda p: (p.n_kutusu.setValue(5), p.yuzde.setChecked(True)),
     dict(tur="ilk_son", n=5, alt=True, yuzde=True)),
    ("ortalama_alti", lambda p: p.std.setValue(2), dict(tur="ortalama", ustu=False, std=2)),
    ("formul", lambda p: p.formul.setText("$A1>10"), dict(tur="formul", formul1="=$A1>10")),
])
def test_kural_penceresi_kural_olusturur(app, tur, ayar, beklenen):
    p = KuralPenceresi(CalismaKitabi(), tur=tur)
    ayar(p)
    kr = p.kural_olustur()
    for alan, deger in beklenen.items():
        assert getattr(kr, alan) == deger, alan
    assert kr.bicim == HAZIR_BICIMLER[0][1] and kr.araliklar == ()


def test_kural_penceresi_renk_olcegi_cubuk_ve_hatalar(app):
    p = KuralPenceresi(CalismaKitabi(), tur="renk_olcegi_3")
    kr = p.kural_olustur()
    assert kr.tur == "renk_olcegi" and [n[0] for n in kr.noktalar] == ["min", "yuzdebirlik", "max"] and kr.bicim is None
    p = KuralPenceresi(CalismaKitabi(), tur="renk_olcegi_2")
    assert len(p.kural_olustur().noktalar) == 2
    p = KuralPenceresi(CalismaKitabi(), tur="veri_cubugu")
    p.cubuk_renk.setCurrentIndex(p.cubuk_renk.findData("kirmizi"))
    p.gradyan.setChecked(False)
    kr = p.kural_olustur()
    assert (kr.renk, kr.gradyan) == ("#F8696B", False)
    for hatali in (KuralPenceresi(CalismaKitabi(), tur="hucre_degeri"), KuralPenceresi(CalismaKitabi(), tur="metin"),
                   KuralPenceresi(CalismaKitabi(), tur="formul")):
        with pytest.raises(ValueError):
            hatali.kural_olustur()


def test_ozel_bicim_yamasi(app):
    pen = OzelBicimPenceresi()
    pen.renkleri_ayarla("#112233", "#FFEEDD")
    pen.kalin.setChecked(True)
    pen.alti_cizili.setChecked(True)
    y = pen.yama()
    assert y == StilYama(yazi_rengi="#112233", dolgu="#FFEEDD", kalin=True, alti_cizili="tek")
    p = KuralPenceresi(CalismaKitabi(), tur="bos")
    p.ozel_bicim_ayarla(y)
    assert p.kural_olustur().bicim == y
    assert OzelBicimPenceresi(yama=y).kalin.isChecked()


def test_yonetici_liste_islemleri(app):
    k = CalismaKitabi()
    s = k.sayfalar[0]
    a = KosulluKural((Aralik(1, 1, 5, 1),), tur="bos", bicim=KIRMIZI)
    b = KosulluKural((Aralik(1, 3, 5, 3),), tur="hucre_degeri", islec="buyuk", formul1="=5", bicim=HAZIR_BICIMLER[1][1])
    s.kosullulari_yaz([a, b])
    y = KuralYoneticisi(k, s, [Aralik(1, 1, 2, 1)], (1, 1))
    assert y.kapsam.currentData() == "secim" and y.liste_kutusu.count() == 1           # yalnız seçimle kesişen
    y.kapsam.setCurrentIndex(1)
    assert y.liste_kutusu.count() == 2
    y.liste_kutusu.setCurrentRow(0)
    y.tasi(1)
    assert y.liste() == [b, a]
    y.tasi(-1)
    assert y.liste() == [a, b]
    y.liste_kutusu.setCurrentRow(1)
    y.dur_kutusu.setChecked(True)
    assert y.liste()[1].dur and y.liste()[1].tur == "hucre_degeri"
    y.sil()
    assert len(y.liste()) == 1
    y.yeni_kural(kural(tur="formul", formul1="=A1>0", bicim=KIRMIZI))
    assert y.liste()[0].tur == "formul" and y.liste()[0].araliklar == (Aralik(1, 1, 2, 1),)
    y.duzenle(kural(tur="hata", bicim=KIRMIZI))
    assert y.liste()[0].tur == "hata" and y.liste()[0].araliklar == (Aralik(1, 1, 2, 1),)


def test_kural_aciklamalari(app):
    k = CalismaKitabi()
    assert kural_aciklamasi(KosulluKural(A, tur="hucre_degeri", islec="arasinda", formul1="=1", formul2="=5"), k) \
        == "Hücre değeri arasında 1 ve 5"
    assert kural_aciklamasi(KosulluKural(A, tur="metin", islec="baslar", metin="x"), k) == "Metin ile başlar: x"
    assert kural_aciklamasi(KosulluKural(A, tur="ilk_son", n=10, yuzde=True), k) == "İlk 10%"
    assert "3 renkli" in kural_aciklamasi(KosulluKural(A, tur="renk_olcegi", noktalar=RENK_OLCEKLERI["yesil_sari_kirmizi"]), k)


# ----------------------------------------------------------------------------- .xlsx
def test_xlsx_gidis_donus_ve_oncelik(tmp_path):
    k = CalismaKitabi()
    s = k.sayfalar[0]
    kr = [
        KosulluKural(A, tur="hucre_degeri", islec="arasinda", formul1="=2", formul2="=$B$1",
                     bicim=StilYama(yazi_rengi="#9C0006", dolgu="#FFC7CE", kalin=True), dur=True),
        KosulluKural(A, tur="metin", islec="icerir", metin='a"b', bicim=StilYama(dolgu="#FFEB9C")),
        KosulluKural(A, tur="bos", bicim=StilYama(kenar={y: ("ince", "#FF0000") for y in ("ust", "alt", "sol", "sag")})),
        KosulluKural(A, tur="yinelenen", benzersiz=True, bicim=StilYama(italik=True)),
        KosulluKural(A, tur="ilk_son", n=3, alt=True, yuzde=True, bicim=StilYama(dolgu="#00FF00")),
        KosulluKural(A, tur="ortalama", ustu=False, esit=True, std=1, bicim=StilYama(dolgu="#00FF00")),
        KosulluKural(A, tur="formul", formul1="=$A1>$B1", bicim=StilYama(dolgu="#112233")),
        KosulluKural(A, tur="renk_olcegi", bicim=None, noktalar=RENK_OLCEKLERI["yesil_sari_kirmizi"]),
        KosulluKural(A, tur="veri_cubugu", bicim=None, renk="#638EC6", yalniz_cubuk=True, max_nokta=("sayi", "=100")),
        KosulluKural((Aralik(1, 3, 2, 3), Aralik(4, 3, 4, 3)), tur="hata_degil", bicim=StilYama(dolgu="#FFFFFF")),
    ]
    s.kosullulari_yaz(kr)
    yol = str(tmp_path / "kb.xlsx")
    assert not xlsx_io.kaydet(k, yol).uyarilar
    k2, rapor = xlsx_io.ac(yol)
    assert k2.sayfalar[0].kosullu == kr and not rapor.uyarilar
    assert xlsx_io.tara(yol).uyarilar == []


def test_xlsx_dxf_fill_bgcolor_ve_oncelik_dosyada(tmp_path):
    import zipfile
    k = CalismaKitabi()
    k.sayfalar[0].kosullulari_yaz([
        KosulluKural(A, tur="bos", bicim=StilYama(dolgu="#FFC7CE")),
        KosulluKural(A, tur="hata", bicim=StilYama(dolgu="#00FF00")),
    ])
    yol = str(tmp_path / "p.xlsx")
    xlsx_io.kaydet(k, yol)
    with zipfile.ZipFile(yol) as z:
        sayfa = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        stiller = z.read("xl/styles.xml").decode("utf-8")
    assert 'priority="1"' in sayfa and 'priority="2"' in sayfa and "containsBlanks" in sayfa
    assert "<bgColor rgb=\"00FFC7CE\"" in stiller or "FFC7CE" in stiller


def test_xlsx_desteklenmeyen_tur_uyari_ve_oncelik_sirasi(tmp_path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, Rule
    from openpyxl.styles import PatternFill
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws.conditional_formatting.add("A1:A5", CellIsRule(operator="greaterThan", formula=["3"],
                                                      fill=PatternFill(bgColor="FFC7CE")))
    ws.conditional_formatting.add("A1:A5", Rule(type="timePeriod", timePeriod="today"))
    ws.conditional_formatting.add("B1:B5", CellIsRule(operator="lessThan", formula=["0"],
                                                      fill=PatternFill(bgColor="00FF00")))
    for i, kuralx in enumerate(ws.conditional_formatting):
        for rl in kuralx.rules:
            rl.priority = 10 - i                                            # ters öncelik: B kuralı en yüksek
    yol = str(tmp_path / "u.xlsx")
    wb.save(yol)
    k, rapor = xlsx_io.ac(yol)
    assert any("1 koşullu biçim kuralı" in u for u in rapor.uyarilar)
    kurallar = k.sayfalar[0].kosullu
    assert [x.islec for x in kurallar] == ["kucuk", "buyuk"]                 # priority küçük = yüksek öncelik
    assert kurallar[1].bicim == StilYama(dolgu="#FFC7CE") and kurallar[0].formul1 == "=0"


# ----------------------------------------------------------------------------- yazdırma
def test_yazdirma_koşullu_biçimi_çizer(pencere):
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QImage, QPainter
    from tabcel.ui import yazdir as yz
    g = pencere.izgara()
    g.sec((1, 1), (8, 1))
    pencere.kosullu_kural_ac(sonuc=kural(tur="hucre_degeri", islec="buyuk", formul1="=50", bicim=KIRMIZI))
    pencere.kosullu_hazir("veri_cubugu", "duz_mavi")

    def resim():
        c = yz.SayfaCizici(g.sayfa, g.yazdirma_yuksekligi, g.font())
        img = QImage(600, 400, QImage.Format_RGB32)
        img.fill(QColor("white"))
        p = QPainter(img)
        c.ciz(p, c.dilimler()[0], QRectF(0, 0, 600, 400), 96.0)
        p.end()
        return img, c
    img, c = resim()
    h = [g.yazdirma_yuksekligi(i) for i in range(1, 9)]
    w = c.genislik(1)
    y6 = sum(h[:5]) + h[5] - 4                                         # satır 6 (60 > 50): kırmızı dolgu, sağ alt köşe
    assert QColor(img.pixel(int(w) - 3, int(y6))).red() > 200 and QColor(img.pixel(int(w) - 3, int(y6))).green() < 60
    y8 = sum(h[:7]) + h[7] / 2                                         # satır 8: en uzun çubuk (mavi)
    px = QColor(img.pixel(int(w * 0.6), int(y8)))
    assert px.blue() > px.red()
    y1 = h[0] / 2                                                       # satır 1: kısa çubuk, orta boş
    assert QColor(img.pixel(int(w * 0.8), int(y1))).red() > 240
    g.sayfa.kosullulari_yaz([])
    img2, _ = resim()
    assert QColor(img2.pixel(int(w) - 3, int(y6))) == QColor("white")
