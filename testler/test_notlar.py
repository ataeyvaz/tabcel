"""Faz 5 / 5e: hücre notları — model, komut, ekle-sil, sırala, taşı, pano, .xlsx (Qt'siz)."""
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from tabcel.io import xlsx_io
from tabcel.model import pano
from tabcel.model.adres import Aralik
from tabcel.model.komutlar import (
    GeriAlYigini, Grup, HucrelerDegis, HucreTasi, NotlarDegis, SatirEkleSil, SutunEkleSil, Sirala, ana_sayfa,
)
from tabcel.model.notlar import Not
from tabcel.model.sayfa import CalismaKitabi, Hucre


@pytest.fixture
def kitap():
    k = CalismaKitabi()
    return k


@pytest.fixture
def sayfa(kitap):
    return kitap.sayfalar[0]


def olaylar(kitap):
    liste = []
    kitap.dinle(liste.append)
    return liste


class TestModel:
    def test_not_degismez_ve_kutu_metni(self):
        n = Not("Kontrol et", "Ata")
        with pytest.raises(Exception):
            n.metin = "x"
        assert n.kutu_metni() == "Ata:\nKontrol et"
        assert Not("Düz").kutu_metni() == "Düz"
        assert Not("", "").bos_mu()

    def test_yaz_ve_eskiler(self, sayfa):
        eski = sayfa.notlari_yaz({(2, 3): Not("a", "Y")})
        assert sayfa.notlar == {(2, 3): Not("a", "Y")}
        assert eski == {(2, 3): None}
        eski2 = sayfa.notlari_yaz({(2, 3): Not("b", "Y")})
        assert eski2 == {(2, 3): Not("a", "Y")}
        sayfa.notlari_yaz(eski2)
        assert sayfa.notlar[(2, 3)].metin == "a"

    def test_none_ve_bos_not_siler(self, sayfa):
        sayfa.notlari_yaz({(1, 1): Not("a"), (2, 2): Not("b")})
        sayfa.notlari_yaz({(1, 1): None, (2, 2): Not("")})
        assert sayfa.notlar == {}

    def test_tek_ek_olayi(self, kitap, sayfa):
        o = olaylar(kitap)
        sayfa.notlari_yaz({(1, 1): Not("a"), (5, 5): Not("b")})
        ek = [x for x in o if x.tur == "ek"]
        assert len(ek) == 1
        assert ek[0]["alan"] == "not" and ek[0]["anahtarlar"] == {(1, 1), (5, 5)}

    def test_degisiklik_yoksa_olay_yok(self, kitap, sayfa):
        sayfa.notlari_yaz({(1, 1): Not("a")})
        o = olaylar(kitap)
        sayfa.notlari_yaz({(1, 1): Not("a"), (9, 9): None})
        assert o == []


class TestKomut:
    def test_yap_geri_al_yinele(self, sayfa):
        y = GeriAlYigini()
        y.calistir(NotlarDegis(sayfa, {(2, 2): Not("x", "A")}))
        assert sayfa.notlar[(2, 2)].metin == "x"
        y.geri_al()
        assert sayfa.notlar == {}
        y.yinele()
        assert sayfa.notlar[(2, 2)].yazar == "A"

    def test_bos_mu(self, sayfa):
        sayfa.notlari_yaz({(1, 1): Not("a")})
        assert NotlarDegis(sayfa, {(1, 1): Not("a")}).bos_mu()
        assert NotlarDegis(sayfa, {(4, 4): None}).bos_mu()
        assert not NotlarDegis(sayfa, {(1, 1): Not("b")}).bos_mu()
        assert not NotlarDegis(sayfa, {(1, 1): None}).bos_mu()

    def test_ana_sayfa(self, sayfa):
        assert ana_sayfa(NotlarDegis(sayfa, {}), True) is sayfa

    def test_hucre_degeri_notu_silmez(self, sayfa):
        sayfa.notlari_yaz({(1, 1): Not("a")})
        y = GeriAlYigini()
        y.calistir(HucrelerDegis(sayfa, {(1, 1): Hucre(5.0)}))
        y.calistir(HucrelerDegis(sayfa, {(1, 1): None}))          # Delete: yalnız içerik (Excel)
        assert (1, 1) in sayfa.notlar


class TestEkleSil:
    def test_satir_ekle_notu_kaydirir_geri_al(self, kitap, sayfa):
        sayfa.notlari_yaz({(3, 2): Not("a")})
        y = GeriAlYigini()
        y.calistir(SatirEkleSil(sayfa, 2, 2, True))
        assert list(sayfa.notlar) == [(5, 2)]
        y.geri_al()
        assert list(sayfa.notlar) == [(3, 2)]

    def test_satir_sil_notu_siler_geri_al_getirir(self, sayfa):
        sayfa.notlari_yaz({(3, 2): Not("a"), (6, 2): Not("b")})
        y = GeriAlYigini()
        y.calistir(SatirEkleSil(sayfa, 3, 1, False))
        assert sayfa.notlar == {(5, 2): Not("b")}
        y.geri_al()
        assert sayfa.notlar == {(3, 2): Not("a"), (6, 2): Not("b")}

    def test_sutun_ekle_sil(self, sayfa):
        sayfa.notlari_yaz({(1, 4): Not("a")})
        y = GeriAlYigini()
        y.calistir(SutunEkleSil(sayfa, 1, 1, True))
        assert list(sayfa.notlar) == [(1, 5)]
        y.calistir(SutunEkleSil(sayfa, 5, 1, False))
        assert sayfa.notlar == {}
        y.geri_al()
        y.geri_al()
        assert list(sayfa.notlar) == [(1, 4)]


class TestSirala:
    def test_not_satirla_birlikte_tasinir(self, sayfa):
        for r, v in enumerate((30.0, 10.0, 20.0), start=1):
            sayfa.ata(r, 1, Hucre(v))
        sayfa.notlari_yaz({(1, 1): Not("otuz"), (2, 1): Not("on")})
        y = GeriAlYigini()
        y.calistir(Sirala(sayfa, Aralik(1, 1, 3, 1), [2, 3, 1]))
        assert [sayfa.deger(r, 1) for r in (1, 2, 3)] == [10.0, 20.0, 30.0]
        assert sayfa.notlar == {(1, 1): Not("on"), (3, 1): Not("otuz")}
        y.geri_al()
        assert sayfa.notlar == {(1, 1): Not("otuz"), (2, 1): Not("on")}


class TestTasi:
    def test_kes_yapistir_notu_tasir_geri_al(self, sayfa):
        sayfa.ata(1, 1, Hucre(1.0))
        sayfa.notlari_yaz({(1, 1): Not("a"), (9, 9): Not("hedefte eski")})
        y = GeriAlYigini()
        y.calistir(HucreTasi(sayfa, (1, 1, 1, 1), sayfa, (9, 9)))
        assert sayfa.notlar == {(9, 9): Not("a")}                   # kaynaktan gitti, hedefin eskisi ezildi
        y.geri_al()
        assert sayfa.notlar == {(1, 1): Not("a"), (9, 9): Not("hedefte eski")}

    def test_baska_sayfaya(self, kitap, sayfa):
        s2 = kitap.sayfa_ekle("S2")
        sayfa.ata(2, 2, Hucre(1.0))
        sayfa.notlari_yaz({(2, 2): Not("a")})
        y = GeriAlYigini()
        y.calistir(HucreTasi(sayfa, (2, 2, 2, 2), s2, (4, 4)))
        assert sayfa.notlar == {} and s2.notlar == {(4, 4): Not("a")}
        y.geri_al()
        assert sayfa.notlar == {(2, 2): Not("a")} and s2.notlar == {}

    def test_notsuz_tasima_notlara_dokunmaz(self, sayfa):
        sayfa.ata(1, 1, Hucre(1.0))
        sayfa.notlari_yaz({(5, 5): Not("uzak")})
        HucreTasi(sayfa, (1, 1, 1, 1), sayfa, (2, 2)).yap()
        assert sayfa.notlar == {(5, 5): Not("uzak")}


class TestPano:
    def test_veri_notu_tasir_json(self, sayfa):
        sayfa.ata(2, 2, Hucre(1.0))
        sayfa.notlari_yaz({(2, 2): Not("a\nb", "Y"), (8, 8): Not("kapsam dışı")})
        v = pano.pano_verisi_al(sayfa, Aralik(2, 2, 3, 3))
        assert v.notlar == {(0, 0): Not("a\nb", "Y")}
        geri = pano.PanoVerisi.json_oku(v.json_yaz())
        assert geri.notlar == {(0, 0): Not("a\nb", "Y")}

    def test_yapistir_notu_getirir(self, sayfa):
        sayfa.ata(1, 1, Hucre(1.0))
        sayfa.notlari_yaz({(1, 1): Not("a", "Y")})
        v = pano.pano_verisi_al(sayfa, Aralik(1, 1, 1, 1))
        plan = pano.yapistirma_notlari(v, sayfa, 4, 4)
        assert plan == {(4, 4): Not("a", "Y")}

    def test_notsuz_kaynak_hedefteki_notu_siler(self, sayfa):
        sayfa.ata(1, 1, Hucre(1.0))
        sayfa.notlari_yaz({(4, 4): Not("eski")})
        v = pano.pano_verisi_al(sayfa, Aralik(1, 1, 1, 1))
        assert pano.yapistirma_notlari(v, sayfa, 4, 4) == {(4, 4): None}

    def test_desen_seciliye_yinelenir(self, sayfa):
        sayfa.ata(1, 1, Hucre(1.0))
        sayfa.notlari_yaz({(1, 1): Not("a")})
        v = pano.pano_verisi_al(sayfa, Aralik(1, 1, 1, 1))
        plan = pano.yapistirma_notlari(v, sayfa, 3, 1, Aralik(3, 1, 5, 1))
        assert plan == {(3, 1): Not("a"), (4, 1): Not("a"), (5, 1): Not("a")}

    def test_dis_veri_ve_notsuz_durum_bos(self, sayfa):
        v = pano.metinden_pano_verisi([["a", "b"]])
        assert pano.yapistirma_notlari(v, sayfa, 1, 1) == {}
        sayfa.ata(1, 1, Hucre(1.0))
        v2 = pano.pano_verisi_al(sayfa, Aralik(1, 1, 1, 1))
        assert pano.yapistirma_notlari(v2, sayfa, 3, 3) == {}      # ne kaynakta ne hedefte not var


class TestXlsx:
    def test_gidis_donus(self, tmp_path):
        k = CalismaKitabi()
        s = k.sayfalar[0]
        s.ata(1, 1, Hucre("x"))
        s.notlari_yaz({(1, 1): Not("Satır1\nSatır2 çğıöşü", "Ata"), (5, 3): Not("boş hücrede", "")})
        yol = str(tmp_path / "n.xlsx")
        xlsx_io.kaydet(k, yol)
        k2, rapor = xlsx_io.ac(yol)
        s2 = k2.sayfalar[0]
        assert s2.notlar[(1, 1)] == Not("Satır1\nSatır2 çğıöşü", "Ata")
        assert s2.notlar[(5, 3)].metin == "boş hücrede"
        assert not rapor.uyarilar                                  # notlar artık veri kaybı sayılmaz

    def test_excel_onekli_metin_yazar_ayrilir(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws["B2"] = 1
        ws["B2"].comment = Comment("Ayşe:\nÖdeme bekliyor", "Ayşe")
        yol = str(tmp_path / "e.xlsx")
        wb.save(yol)
        k, rapor = xlsx_io.ac(yol)
        assert k.sayfalar[0].notlar[(2, 2)] == Not("Ödeme bekliyor", "Ayşe")
        assert not rapor.uyarilar

    def test_onek_yoksa_metin_oldugu_gibi(self, tmp_path):
        wb = Workbook()
        wb.active["A1"].comment = Comment("Sadece metin", "Veli")
        yol = str(tmp_path / "o.xlsx")
        wb.save(yol)
        k, _ = xlsx_io.ac(yol)
        assert k.sayfalar[0].notlar[(1, 1)] == Not("Sadece metin", "Veli")

    def test_dizili_aciklama_uyari(self, tmp_path):
        import zipfile
        wb = Workbook()
        wb.active["A1"].comment = Comment("[Threaded comment]\n\nBir yorum", "X")
        yol = str(tmp_path / "t.xlsx")
        wb.save(yol)
        with zipfile.ZipFile(yol, "a") as z:                        # dizili açıklama parçasının varlığını taklit et
            z.writestr("xl/threadedComments/threadedComment1.xml", "<ThreadedComments/>")
        k, rapor = xlsx_io.ac(yol)
        assert (1, 1) in k.sayfalar[0].notlar
        assert rapor.veri_kaybi and any("dizili" in u for u in rapor.uyarilar)

    def test_notsuz_dosya_notsuz(self, tmp_path):
        k = CalismaKitabi()
        k.sayfalar[0].ata(1, 1, Hucre(1.0))
        yol = str(tmp_path / "s.xlsx")
        xlsx_io.kaydet(k, yol)
        assert load_workbook(yol).active["A1"].comment is None
        assert xlsx_io.ac(yol)[0].sayfalar[0].notlar == {}
