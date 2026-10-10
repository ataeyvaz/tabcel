"""Faz 4 / 4g: `testler/ornekler/` gerçek örnek dosyalarıyla gidiş-dönüş ve uçtan uca iş akışı."""
import os
import re
import shutil
import zipfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from tabcel.formul import ceviri
from tabcel.formul.motor import Hesaplayici
from tabcel.io import csv_io, xlsx_io
from tabcel.model.adres import Aralik
from tabcel.model.deger import BOL0
from tabcel.model.notlar import Not
from tabcel.model.sayfa_yapisi import SayfaYapisi
from tabcel.ui.dosya_islemleri import DosyaIslemleri
from tabcel.ui.window import Pencere

ORNEKLER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ornekler")


def ornek(ad):
    return os.path.join(ORNEKLER, ad)


def kitap_farklari(a, b, tolerans_px=8):
    """İki kitap arasındaki farkların listesi (boşsa eşit): değer/formül, stil, ölçü, birleşim, görünüm."""
    f = []
    if [s.ad for s in a.sayfalar] != [s.ad for s in b.sayfalar]:
        return ["sayfa adları/sırası"]
    for s1, s2 in zip(a.sayfalar, b.sayfalar):
        ad = s1.ad
        if set(s1._hucreler) != set(s2._hucreler):
            f.append(f"{ad}: hücre kümesi")
        for k, h in s1._hucreler.items():
            h2 = s2._hucreler.get(k)
            if h2 is None:
                continue
            if h.formul is not None:
                if ceviri.disa(h.formul, a) != ceviri.disa(h2.formul or "", b):
                    f.append(f"{ad}!{k}: formül")
            elif h != h2:
                f.append(f"{ad}!{k}: değer {h.deger!r} → {h2.deger!r}")
        for k in set(s1._hucreler) | set(s1._stil) | set(s2._stil):
            if s1.stil(*k) != s2.stil(*k):
                f.append(f"{ad}!{k}: stil")
        for sat in set(s1.satir_stil) | set(s2.satir_stil):
            if s1.havuz.stil(s1.satir_stil.get(sat, 0)) != s2.havuz.stil(s2.satir_stil.get(sat, 0)):
                f.append(f"{ad}: satır {sat} stili")
        for c in set(s1.sutun_stil) | set(s2.sutun_stil):
            if s1.havuz.stil(s1.sutun_stil.get(c, 0)) != s2.havuz.stil(s2.sutun_stil.get(c, 0)):
                f.append(f"{ad}: sütun {c} stili")
        for c in set(s1.sutun_genislik) | set(s2.sutun_genislik):
            if abs(s1.genislik(c) - s2.genislik(c)) > tolerans_px:
                f.append(f"{ad}: sütun {c} genişliği")
        for r in set(s1.satir_yukseklik) | set(s2.satir_yukseklik):
            if abs(s1.yukseklik(r) - s2.yukseklik(r)) > 3:
                f.append(f"{ad}: satır {r} yüksekliği")
        if (s1.gizli_sutun, s1.gizli_satir) != (s2.gizli_sutun, s2.gizli_satir):
            f.append(f"{ad}: gizli satır/sütun")
        if sorted(s1.birlesimler, key=str) != sorted(s2.birlesimler, key=str):
            f.append(f"{ad}: birleşim")
        if (s1.dondur, s1.kilavuz, s1.sekme_rengi, s1.sayfa_yapisi) != (s2.dondur, s2.kilavuz, s2.sekme_rengi,
                                                                       s2.sayfa_yapisi):
            f.append(f"{ad}: görünüm/sayfa yapısı")
    return f


# ---------------------------------------------------------------- butce.xlsx
def test_butce_okunur_beklenen_icerik():
    k, rapor = xlsx_io.ac(ornek("butce.xlsx"))
    s, ö = k.sayfalar
    assert (s.ad, ö.ad) == ("Bütçe", "Özet") and not rapor.veri_kaybi and rapor.uyarilar == []
    assert s.birlesimler == [Aralik(1, 1, 1, 6)] and s.stil(1, 1).kalin and s.stil(1, 1).boyut == 16.0
    assert s.stil(1, 1).dolgu == "#217346" and s.stil(1, 1).yatay == "orta" and s.stil(1, 1).yazi_rengi == "#FFFFFF"
    assert s.deger(3, 3) == 42000.0 and s.stil(3, 3).sayi_bicimi == '#,##0.00 "₺"'
    assert ceviri.disa(s.hucre(3, 6).formul, k) == "=SUM(C3:E3)"
    assert ceviri.disa(s.hucre(9, 6).formul, k) == '=SUMIF(B3:B7,"Gider",F3:F7)'
    assert s.deger(11, 6) == 46304.0 and s.stil(11, 6).sayi_bicimi == "dd.mm.yyyy"      # 09.10.2026
    assert (s.dondur, s.kilavuz, s.sekme_rengi) == ((3, 3), False, "#217346")
    assert 8 in s.gizli_sutun and 12 in s.gizli_satir
    assert s.satir_yukseklik[1] > 30
    y = s.sayfa_yapisi
    assert (y.yon, y.kagit, y.baslik_satirlari, y.kilavuz) == ("yatay", "A4", (1, 2), True)
    assert s.stil(13, 2).kaydir and s.stil(13, 2).dikey == "ust"
    # dosyada 'Bütçe'!F9 (tırnaklı) idi; Excel'in kendi yazımı da tırnaksızdır: eşdeğer, bizimki tırnaksız yazar
    assert ceviri.disa(ö.hucre(2, 2).formul, k) == "=Bütçe!F9"
    assert ceviri.disa(ö.hucre(3, 2).formul, k) == "=Bütçe!C3"
    assert ö.deger(4, 2) is True


def test_butce_hesap_dogru():
    k, _ = xlsx_io.ac(ornek("butce.xlsx"))
    Hesaplayici(k)
    s, ö = k.sayfalar
    assert s.sonuc(3, 6) == 129000.0
    gider = sum(s.sonuc(r, 6) for r in range(4, 8))
    assert s.sonuc(9, 6) == pytest.approx(gider)
    assert s.sonuc(10, 6) == pytest.approx(gider / 129000.0) and s.sonuc(12, 6) == "Dikkat"     # oran %63,7 > %50
    assert ö.sonuc(2, 2) == pytest.approx(gider) and ö.sonuc(3, 2) == 42000.0 and ö.sonuc(5, 2) == BOL0


def test_butce_gidis_donus_kayipsiz(tmp_path):
    k1, _ = xlsx_io.ac(ornek("butce.xlsx"))
    yol = str(tmp_path / "kopya.xlsx")
    xlsx_io.kaydet(k1, yol)
    k2, rapor = xlsx_io.ac(yol)
    assert kitap_farklari(k1, k2) == [] and not rapor.veri_kaybi
    # ikinci tur da kararlı (kaydet → aç → kaydet → aç)
    yol2 = str(tmp_path / "kopya2.xlsx")
    xlsx_io.kaydet(k2, yol2)
    assert kitap_farklari(k2, xlsx_io.ac(yol2)[0]) == []


def test_butce_openpyxl_ile_de_ayni(tmp_path):
    """Bizim yazdığımız dosyayı başka bir okuyucu (openpyxl) de aynı görür."""
    from openpyxl import load_workbook
    k, _ = xlsx_io.ac(ornek("butce.xlsx"))
    yol = str(tmp_path / "k.xlsx")
    xlsx_io.kaydet(k, yol)
    a, b = load_workbook(ornek("butce.xlsx")), load_workbook(yol)
    for ad in a.sheetnames:
        for satir in a[ad].iter_rows():
            for h in satir:
                h2 = b[ad][h.coordinate]
                v1 = h.value.replace("'", "") if isinstance(h.value, str) and h.value.startswith("=") else h.value
                assert v1 == h2.value and h.number_format == h2.number_format, (ad, h.coordinate)
                assert h.font.b == h2.font.b and h.fill.fgColor.rgb == h2.fill.fgColor.rgb
    assert b["Bütçe"].freeze_panes == "C3" and str(b["Bütçe"].merged_cells) == "A1:F1"
    assert b["Bütçe"].column_dimensions["H"].hidden and b["Bütçe"].print_title_rows == "$1:$2"


# ---------------------------------------------------------------- grafikli.xlsx
def test_grafikli_uyarilar_ve_veri_korunur(tmp_path):
    k, rapor = xlsx_io.ac(ornek("grafikli.xlsx"))
    assert rapor.veri_kaybi
    metin = " ".join(rapor.uyarilar)
    for ne in ("grafik", "resim", "koşullu biçim", "veri doğrulama", "filtre"):
        assert ne in metin, ne
    assert "tanımlı ad" not in metin                       # 5d: tanımlı adlar artık korunur
    assert "yorum" not in metin                            # 5e: notlar artık korunur (uyarı yok)
    assert k.ad_bul("SatisToplam") is not None             # Faz 5 / 5d
    s = k.sayfalar[0]
    assert s.notlar == {(2, 2): Not("İlk ay değeri", "Ata")}   # Faz 5 / 5e
    assert s.deger(2, 1) == "Ocak" and s.deger(5, 2) == 40.0 and s.deger(2, 4) == "Evet"
    # kaydedilince korunamayan öğeler çıkar (uyarılmıştı), veri kalır
    yol = str(tmp_path / "g.xlsx")
    xlsx_io.kaydet(k, yol)
    with zipfile.ZipFile(yol) as z:
        adlar = " ".join(z.namelist())
    assert "charts" not in adlar and "media" not in adlar
    k2, r2 = xlsx_io.ac(yol)
    assert not r2.veri_kaybi and kitap_farklari(k, k2) == []


def test_ozgun_dosya_degismez(tmp_path):
    kopya = str(tmp_path / "g.xlsx")
    shutil.copy(ornek("grafikli.xlsx"), kopya)
    onceki = open(kopya, "rb").read()
    k, _ = xlsx_io.ac(kopya)                                    # açmak dosyaya yazmaz
    assert open(kopya, "rb").read() == onceki


# ---------------------------------------------------------------- CSV
def test_turkce_csv():
    satirlar, kod, ayrac = csv_io.dosya_oku(ornek("satis_tr.csv"))
    assert (kod, ayrac) == ("cp1254", ";") and satirlar[0] == ["Ürün", "Adet", "Fiyat", "Tarih"]
    k, uyari = csv_io.kitap_olustur(satirlar, "satis_tr")
    s = k.sayfalar[0]
    assert uyari == [] and s.deger(2, 1) == "Çay" and s.deger(2, 3) == 12.5 and s.deger(3, 3) == 8.25
    assert s.deger(2, 4) == 46307.0 and s.stil(2, 4).sayi_bicimi != "General" and s.deger(4, 1) == "İncir"


def test_utf8_virgullu_csv_tirnakli_alan():
    satirlar, kod, ayrac = csv_io.dosya_oku(ornek("satis_utf8.csv"))
    assert (kod, ayrac) == ("utf-8", ",")
    assert satirlar == [["Ad", "Not"], ["Ali, Veli", "merhaba"], ["Ayşe", "iki\r\nsatır"]]


# ---------------------------------------------------------------- uçtan uca iş akışı
@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def _sessiz(app, monkeypatch):
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: QMessageBox.Ok)
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: QMessageBox.Ok)
    yield
    for p in list(DosyaIslemleri._tutulan):
        p.degisti_mi = lambda: False
        p.close()
    DosyaIslemleri._tutulan.clear()


def test_is_akisi_ac_duzenle_kaydet_pdf_csv(tmp_path, monkeypatch):
    kopya = str(tmp_path / "butce.xlsx")
    shutil.copy(ornek("butce.xlsx"), kopya)
    ilk = Pencere(xlsx_io.ac(kopya)[0])
    ilk.show()
    p = ilk.dosya_ac(kopya)                                     # diskten açılış (yeni pencere)
    g = p.izgara()
    assert p.windowTitle() == "butce.xlsx - Tabcel" and not p.degisti_mi()
    assert g.sayfa.sonuc(3, 6) == 129000.0
    # düzenleme: Mart maaşı değişir → formüller yeniden hesaplanır, belge "değişti"
    g.hucre_gir(3, 5, "50000")
    assert g.sayfa.sonuc(3, 6) == 134000.0 and p.degisti_mi() and p.windowTitle().startswith("butce.xlsx *")
    assert p.otokayit_yap()                                     # kurtarma kopyası
    assert p.kaydet() and not p.degisti_mi()                    # yol var: diyalogsuz üzerine yazar
    # yeniden aç: yeni değer dosyada
    p.yol = None
    q = p.dosya_ac(kopya)
    assert q.izgara().sayfa.sonuc(3, 6) == 134000.0
    # PDF: tüm sayfalar (Bütçe yatay + Özet dikey) ve CSV (görünen metin, etkin sayfa)
    pdf = str(tmp_path / "butce.pdf")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (pdf, ""))
    assert q.pdf_aktar(tum=True)
    veri = open(pdf, "rb").read()
    assert veri.startswith(b"%PDF")
    assert len(re.findall(rb"/Type\s*/Page(?![s\w])", veri)) >= 2
    kutular = [tuple(round(float(x)) for x in m) for m in re.findall(
        rb"/MediaBox\s*\[\s*([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s*\]", veri)]
    assert (0, 0, 842, 595) in kutular and (0, 0, 595, 842) in kutular
    csvyol = str(tmp_path / "butce.csv")
    csv_io.dosya_yaz(q.izgara().sayfa, csvyol)
    metin = open(csvyol, encoding="utf-8-sig", newline="").read()
    assert metin.splitlines()[0].startswith("Aylık Bütçe 2026") and "50.000,00 ₺" in metin
    assert "09.10.2026" in metin and "Dikkat" in metin and "61,3%" in metin       # 82.201,5 / 134.000
