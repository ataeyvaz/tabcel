"""Faz 10c: şablon bilgisinin kitapta saklanması (`_tabcel` çok gizli sayfa), sağlama, açılış denetimi, Excel'de düzenleme senaryoları."""
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tabcel.formul.motor import Hesaplayici
from tabcel.io import xlsx_io
from tabcel.model.komutlar import GeriAlYigini, SutunEkleSil
from tabcel.model.sayfa import Hucre
from tabcel.sablon import saklama, yukleyici
from tabcel.sablon.uretici import uret


@pytest.fixture(scope="module")
def sablon():
    return yukleyici.depo_yukle(yukleyici.yerlesik_klasor(), Path("yok_klasor")).bul("yerlesik:stok_takibi")


def kaydet_ac(kitap, tmp_path, ad="k.xlsx"):
    yol = str(tmp_path / ad)
    xlsx_io.kaydet(kitap, yol)
    return yol, *xlsx_io.ac(yol)


def sayfa_(k, ad):
    return next(s for s in k.sayfalar if s.ad == ad)


# ---------------------------------------------------------------- kayıt ve serileştirme
def test_uretim_kayit_koyar(sablon):
    r = uret(sablon, {"depo": True}, ornek_veri=True)
    k = r.kitap.sablon
    assert k.akis_kimligi == "stok_takibi" and k.cevaplar["depo"] is True and k.uretici.startswith("Tabcel")
    assert k.baglar["urunler"].basliklar["urun_kodu"] == "Ürün kodu" and k.baglar["urunler"].sayfa_adi == "Ürünler"
    assert set(k.baglar) == {"urunler", "stok_hareketleri", "stok_durumu", "depolar"}
    assert list(k.baglar["stok_durumu"].hesaplanan_ozet) == ["urun_kodu", "giris", "cikis", "mevcut", "durum"]
    assert k.ornek_satirlar == {"urunler": [2, 3], "stok_hareketleri": [2, 3, 4]}
    assert k.ham_akis["kimlik"] == "stok_takibi" and set(k.ham_moduller) == {m for m, _ in sablon.akis.moduller}


def test_json_gidis_donus_ve_parcalama(sablon, monkeypatch):
    k = uret(sablon).kitap.sablon
    metin = saklama.json_yaz(k)
    assert saklama.json_yaz(saklama.json_oku(metin)) == metin
    monkeypatch.setattr(saklama, "PARCA_BOYU", 997)
    satirlar = saklama.satirlar(k)
    assert satirlar[0][0] == saklama.ETIKET and satirlar[0][3] == len(satirlar) - 1 > 5
    assert all(len(s[0]) <= 997 for s in satirlar[1:])
    k2, neden = saklama.coz(satirlar)
    assert neden is None and saklama.json_yaz(k2) == metin


def test_gercek_boyutta_parca_siniri(sablon):
    """Excel hücre sınırı 32.767; parçalar 30.000'i aşmaz."""
    k = uret(sablon).kitap.sablon
    for s in saklama.satirlar(k)[1:]:
        assert len(s[0]) <= saklama.PARCA_BOYU < 32767


# ---------------------------------------------------------------- dosya
def test_xlsx_cok_gizli_sayfa_ve_model(sablon, tmp_path):
    r = uret(sablon, ornek_veri=True)
    yol, k2, rapor = kaydet_ac(r.kitap, tmp_path)
    wb = load_workbook(yol)
    assert wb.sheetnames[-1] == "_tabcel" and wb["_tabcel"].sheet_state == "veryHidden"
    assert wb["_tabcel"]["A1"].value == saklama.ETIKET
    assert "_tabcel" not in [s.ad for s in k2.sayfalar]                              # modelde sıradan sayfa değil
    assert [s.ad for s in k2.sayfalar] == [s.ad for s in r.kitap.sayfalar]
    assert k2.sablon is not None and k2.sablon.akis_kimligi == "stok_takibi"
    assert k2.sablon_denetim.aktif and k2.sablon_denetim.bulgular == [] and k2.sablon_ham is None
    assert not rapor.veri_kaybi and not any("Şablon" in u for u in rapor.uyarilar)
    assert saklama.json_yaz(k2.sablon) == saklama.json_yaz(r.kitap.sablon)
    assert k2.sablon_denetim.esleme["urunler"]["alis_fiyati"] == 5


def test_sablonsuz_kitapta_sayfa_yok(tmp_path):
    from tabcel.model.sayfa import CalismaKitabi
    k = CalismaKitabi()
    k.sayfalar[0].ata(1, 1, Hucre(1.0))
    yol, k2, _ = kaydet_ac(k, tmp_path)
    assert load_workbook(yol).sheetnames == ["Sayfa1"] and k2.sablon is None and k2.sablon_denetim is None


def test_ikinci_kayit_gidis_donus(sablon, tmp_path):
    r = uret(sablon)
    _, k2, _ = kaydet_ac(r.kitap, tmp_path, "a.xlsx")
    _, k3, _ = kaydet_ac(k2, tmp_path, "b.xlsx")
    assert saklama.json_yaz(k3.sablon) == saklama.json_yaz(r.kitap.sablon) and k3.sablon_denetim.aktif


def test_formul_hucresi_olarak_parca_bozulmaz(sablon, tmp_path, monkeypatch):
    """"=" ile başlayan JSON parçası formül sayılmamalı."""
    r = uret(sablon)
    metin = saklama.json_yaz(r.kitap.sablon)
    monkeypatch.setattr(saklama, "PARCA_BOYU", metin.index("=") if "=" in metin else 50)
    # '=' ile başlayan bir parça oluşana dek sınırı kaydır
    for n in range(40, 400):
        monkeypatch.setattr(saklama, "PARCA_BOYU", n)
        if any(s[0].startswith("=") for s in saklama.satirlar(r.kitap.sablon)[1:]):
            break
    else:
        pytest.skip("'=' ile başlayan parça üretilemedi")
    _, k2, _ = kaydet_ac(r.kitap, tmp_path)
    assert k2.sablon is not None and saklama.json_yaz(k2.sablon) == metin


# ---------------------------------------------------------------- bozulma / sürüm
def _tabcel_duzenle(yol, f):
    wb = load_workbook(yol)
    f(wb["_tabcel"])
    wb.save(yol)


def test_saglama_tutmazsa_form_kapali_veri_korunur(sablon, tmp_path):
    yol, _, _ = kaydet_ac(uret(sablon).kitap, tmp_path)
    _tabcel_duzenle(yol, lambda ws: ws.cell(2, 1).__setattr__("value", ws.cell(2, 1).value.replace("Stok takibi", "Stok TAKIBI", 1)))
    k, rapor = xlsx_io.ac(yol)
    assert k.sablon is None and k.sablon_ham is not None and not k.sablon_denetim.aktif
    assert "sağlama tutmuyor" in k.sablon_denetim.neden
    assert any("Şablon formu kapalı" in u and "sağlama tutmuyor" in u for u in rapor.uyarilar)
    d = sayfa_(k, "Stok Durumu")                                                     # tablo ve formüller çalışmaya devam eder
    assert d._hucreler[(2, 2)].formul is not None
    yol2 = str(tmp_path / "geri.xlsx")
    xlsx_io.kaydet(k, yol2)                                                          # bilgi silinmez, olduğu gibi geri yazılır
    a, b = load_workbook(yol)["_tabcel"], load_workbook(yol2)["_tabcel"]
    assert [[c.value for c in r] for r in a.iter_rows()] == [[c.value for c in r] for r in b.iter_rows()]


def _yeniden_sagla(ws, f):
    """Satırlardaki JSON'u `f(dict)` ile değiştirip sağlamayı yeniden hesaplar (Excel'de elle değil, bozuk-ama-tutarlı durum)."""
    metin = "".join(ws.cell(i, 1).value for i in range(2, 2 + ws.cell(1, 4).value))
    d = f(json.loads(metin))
    yeni = json.dumps(d, ensure_ascii=False, separators=(",", ":"))
    ws.delete_rows(2, ws.max_row)
    ws.cell(1, 4).value = 1
    ws.cell(1, 5).value = saklama.ozet(yeni)
    ws.cell(2, 1).value = yeni


def test_daha_yeni_saklama_surumu(sablon, tmp_path):
    yol, _, _ = kaydet_ac(uret(sablon).kitap, tmp_path)
    _tabcel_duzenle(yol, lambda ws: _yeniden_sagla(ws, lambda d: {**d, "surum": saklama.SURUM + 1}))
    k, _ = xlsx_io.ac(yol)
    assert k.sablon is None and "daha yeni" in k.sablon_denetim.neden and k.sablon_ham


def test_bozuk_json_ve_eksik_alan(sablon, tmp_path):
    yol, _, _ = kaydet_ac(uret(sablon).kitap, tmp_path)
    _tabcel_duzenle(yol, lambda ws: _yeniden_sagla(ws, lambda d: {"surum": 1}))
    k, _ = xlsx_io.ac(yol)
    assert k.sablon is None and "eksik alan" in k.sablon_denetim.neden
    # parçalar eksik
    yol, _, _ = kaydet_ac(uret(sablon).kitap, tmp_path, "e.xlsx")
    _tabcel_duzenle(yol, lambda ws: ws.cell(1, 4).__setattr__("value", 99))
    k, _ = xlsx_io.ac(yol)
    assert k.sablon is None and "parçaları eksik" in k.sablon_denetim.neden
    # etiketsiz sayfa
    yol, _, _ = kaydet_ac(uret(sablon).kitap, tmp_path, "t.xlsx")
    _tabcel_duzenle(yol, lambda ws: ws.cell(1, 1).__setattr__("value", "BASKA"))
    assert "tanınmayan" in xlsx_io.ac(yol)[0].sablon_denetim.neden


def test_tanimi_bozuk(sablon, tmp_path):
    def boz(d):
        d["tanim"]["akis"]["kimlik"] = "BÜYÜK"
        return d
    yol, _, _ = kaydet_ac(uret(sablon).kitap, tmp_path)
    _tabcel_duzenle(yol, lambda ws: _yeniden_sagla(ws, boz))
    k, _ = xlsx_io.ac(yol)
    assert k.sablon is not None and not k.sablon_denetim.aktif and k.sablon_denetim.bulgular[0].tur == "tanim_bozuk"


def test_eski_surumlu_tanim_goc_ile_okunur(sablon, tmp_path, monkeypatch):
    """Kitapta gömülü tanım v1; uygulama şeması v2'ye çıkmış → göç zinciri çalışır, form açık kalır."""
    from tabcel.sablon import sema
    yol, _, _ = kaydet_ac(uret(sablon).kitap, tmp_path)
    monkeypatch.setattr(sema, "SURUM", 2)
    sirayla = []
    monkeypatch.setitem(yukleyici.GOCLER["akis"], 1, lambda d: sirayla.append("akis") or d)
    monkeypatch.setitem(yukleyici.GOCLER["modul"], 1, lambda d: sirayla.append("modul") or d)
    k, _ = xlsx_io.ac(yol)
    assert k.sablon_denetim.aktif and sirayla.count("akis") == 1 and sirayla.count("modul") == 6


# ---------------------------------------------------------------- Excel'de düzenleme senaryoları (uyuşma denetimi)
@pytest.fixture
def acik(sablon, tmp_path):
    _, k, _ = kaydet_ac(uret(sablon, {"depo": True}, ornek_veri=True).kitap, tmp_path)
    Hesaplayici(k)
    return k


def test_sayfa_yeniden_adlandirma_bozmaz(acik):
    acik.yeniden_adlandir(sayfa_(acik, "Ürünler"), "Mallarım")
    d = saklama.denetle(acik, acik.sablon)
    assert d.aktif and d.bulgular == [] and d.sayfalar["urunler"] == sayfa_(acik, "Mallarım").kimlik


def test_sutun_ekleme_baslik_adiyla_yeniden_eslenir(acik):
    urun = sayfa_(acik, "Ürünler")
    GeriAlYigini().calistir(SutunEkleSil(urun, 3, 1, True))                    # Ürün adı ile Barkod arasına
    urun.ata(1, 3, Hucre("Raf"))
    d = saklama.denetle(acik, acik.sablon)
    assert d.aktif
    assert d.esleme["urunler"]["barkod"] == 4 and d.esleme["urunler"]["alis_fiyati"] == 6 and d.esleme["urunler"]["urun_kodu"] == 1
    assert {b.tur for b in d.bulgular} == {"baslik_yer", "yeni_sutun"} and all(b.onem == saklama.BILGI for b in d.bulgular)


def test_sutun_silme_formu_kapatir_tablo_calisir(acik):
    urun = sayfa_(acik, "Ürünler")
    GeriAlYigini().calistir(SutunEkleSil(urun, 3, 1, False))                   # Barkod silindi
    d = saklama.denetle(acik, acik.sablon)
    assert not d.aktif and [b.tur for b in d.bulgular] == ["baslik_yok"] and "“Barkod”" in d.bulgular[0].mesaj
    assert "urunler" not in d.esleme
    sd = sayfa_(acik, "Stok Durumu")
    assert sd.deger(2, 4) == 42.0                                                # tablo hesapları sürer


def test_baslik_yeniden_adlandirma(acik):
    urun = sayfa_(acik, "Ürünler")
    urun.ata(1, 2, Hucre("Malın adı"))
    d = saklama.denetle(acik, acik.sablon)
    assert not d.aktif and "“Ürün adı”" in d.bulgular[0].mesaj


def test_capa_silinirse_sayfa_adiyla_bulunur(acik):
    anahtar = acik.ad_bul("TR_MOD_urunler").anahtar()
    acik.ad_yaz(None, anahtar)
    d = saklama.denetle(acik, acik.sablon)
    assert d.aktif and [b.tur for b in d.bulgular] == ["capa_yok"] and d.esleme["urunler"]["urun_kodu"] == 1


def test_capa_ve_sayfa_yoksa_sayfa_yok(acik):
    acik.ad_yaz(None, acik.ad_bul("TR_MOD_depolar").anahtar())
    acik.yeniden_adlandir(sayfa_(acik, "Depolar"), "Şubeler")
    d = saklama.denetle(acik, acik.sablon)
    assert not d.aktif and [b.tur for b in d.bulgular] == ["sayfa_yok"] and "“Depolar”" in d.bulgular[0].mesaj


def test_formul_degisikligi_bilgi_duzeyinde(acik):
    sd = sayfa_(acik, "Stok Durumu")
    sd.ata(2, 4, Hucre(formul="=$B2-$C2+1"))
    d = saklama.denetle(acik, acik.sablon)
    assert d.aktif and [b.tur for b in d.bulgular] == ["formul_degisti"] and "“Mevcut”" in d.bulgular[0].mesaj


def test_formu_kapat_bilgi_silinir(sablon, tmp_path):
    k = uret(sablon).kitap
    k.sablon, k.sablon_ham = None, None                                           # "Formu kapat (şablon bilgisini sil)"
    yol, k2, _ = kaydet_ac(k, tmp_path)
    assert "_tabcel" not in load_workbook(yol).sheetnames and k2.sablon is None


def test_uyusmazlik_dosyada_kaydedilince_bilgi_korunur(acik, tmp_path):
    """Form devre dışıyken kaydedilen dosyada bilgi silinmez (yeniden eşleme sonra yapılabilir)."""
    GeriAlYigini().calistir(SutunEkleSil(sayfa_(acik, "Ürünler"), 3, 1, False))
    yol, k2, rapor = kaydet_ac(acik, tmp_path, "u.xlsx")
    assert k2.sablon is not None and not k2.sablon_denetim.aktif
    assert any("Şablon formu kapalı" in u and "“Barkod”" in u for u in rapor.uyarilar)


def test_kullanici_sayfasi_ayni_adda(sablon, tmp_path):
    r = uret(sablon)
    r.kitap.sayfa_ekle("_tabcel")
    yol = str(tmp_path / "c.xlsx")
    rapor = xlsx_io.kaydet(r.kitap, yol)
    assert any("zaten var" in u for u in rapor.uyarilar) and rapor.veri_kaybi
    assert load_workbook(yol).sheetnames.count("_tabcel") == 1
