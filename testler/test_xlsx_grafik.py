"""Faz 6 / 6h: grafik ve resmin .xlsx'e yazılması/okunması, Tabcel imi, uyarı mantığı."""
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, BarChart3D, RadarChart, Reference, Series
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from PIL import Image as PilImage

from tabcel.formul.motor import Hesaplayici
from tabcel.io import xlsx_io
from tabcel.model.adres import Aralik
from tabcel.model.grafik import Grafik, Konum, Referans, Resim, Seri
from tabcel.model.grafik_islem import GALERI
from tabcel.model.sayfa import CalismaKitabi, Hucre


def png(boyut=(120, 60), renk=(200, 30, 30)) -> bytes:
    b = io.BytesIO()
    PilImage.new("RGB", boyut, renk).save(b, "PNG")
    return b.getvalue()


@pytest.fixture
def kitap():
    k = CalismaKitabi()
    Hesaplayici(k)
    s = k.sayfalar[0]
    for i, satir in enumerate([[None, "Ocak", "Şubat"], ["Elma", 1, 2], ["Armut", 3, 4], ["Kiraz", 5, 6]], 1):
        for j, v in enumerate(satir, 1):
            if v is not None:
                s.ata(i, j, Hucre(v))
    return k


def ref(s, r1, c1, r2, c2):
    return Referans(s.kimlik, Aralik(r1, c1, r2, c2))


def tam_grafik(s, tur="sutun", alt="kumelenmis", **kw):
    seriler = (Seri(ad_ref=ref(s, 1, 2, 1, 2), deger=ref(s, 2, 2, 4, 2), kategori=ref(s, 2, 1, 4, 1)),
               Seri(ad="İkinci", deger=ref(s, 2, 3, 4, 3), kategori=ref(s, 2, 1, 4, 1), renk="#FF0000"))
    if tur in ("pasta", "halka"):
        seriler = seriler[:1]
    return Grafik(tur=tur, alt_tur=alt, seriler=seriler, konum=Konum(3, 5, 7, 9), boyut=(400, 250), **kw)


def gidis_donus(kitap, yol):
    xlsx_io.kaydet(kitap, str(yol))
    return xlsx_io.ac(str(yol))


def esit(a, b):
    """Kimlik hariç eşitlik (okuyunca kimlikler yeniden atanır)."""
    if isinstance(a, Grafik):
        return a.degistir(kimlik=0) == b.degistir(kimlik=0)
    return a.degistir(kimlik=0, ad="") == b.degistir(kimlik=0, ad="")


# ------------------------------------------------------------------ gidiş-dönüş
@pytest.mark.parametrize("tur,alt", [(t, a) for t, o in GALERI.items() for a, _ in o])
def test_her_tur_ve_alt_tur_gidis_donus(kitap, tmp_path, tur, alt):
    s = kitap.sayfalar[0]
    g = tam_grafik(s, tur, alt, baslik="Başlık", gosterge="alt", x_baslik="Ay", y_baslik="Adet", deger_etiketi=True)
    s.nesne_ekle(g)
    k2, rapor = gidis_donus(kitap, tmp_path / "a.xlsx")
    assert rapor.uyarilar == [] and not rapor.veri_kaybi
    g2 = k2.sayfalar[0].nesneler[0]
    if tur in ("pasta", "halka"):                            # eksen başlığı pastada yoktur
        g = g.degistir(x_baslik=None, y_baslik=None)
    assert esit(g, g2)


def test_varsayilanlar_gidis_donus_ve_baslik_gizli(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s.nesne_ekle(tam_grafik(s))                              # başlıksız, gösterge sağ, renkler palet
    xlsx_io.kaydet(kitap, str(tmp_path / "a.xlsx"))
    with zipfile.ZipFile(tmp_path / "a.xlsx") as z:
        xml = z.read("xl/charts/chart1.xml").decode()
    assert "autoTitleDeleted" in xml                         # Excel tek seride seri adını başlık yapmasın
    k2, _ = xlsx_io.ac(str(tmp_path / "a.xlsx"))
    g2 = k2.sayfalar[0].nesneler[0]
    assert g2.baslik is None and g2.gosterge == "sag" and g2.seriler[0].renk is None
    assert g2.seriler[1].renk == "#FF0000"                   # açık renk korunur


def test_gosterge_yok_ve_eksen_sinirlari(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s.nesne_ekle(tam_grafik(s, "cizgi", gosterge=None, y_min=0.5, y_max=12))
    g2 = gidis_donus(kitap, tmp_path / "a.xlsx")[0].sayfalar[0].nesneler[0]
    assert g2.gosterge is None and (g2.y_min, g2.y_max) == (0.5, 12)


def test_dagilim_x_ve_y_sinirlari(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s.nesne_ekle(tam_grafik(s, "dagilim", x_min=0, x_max=10, y_min=1, y_max=9))
    g2 = gidis_donus(kitap, tmp_path / "a.xlsx")[0].sayfalar[0].nesneler[0]
    assert (g2.x_min, g2.x_max, g2.y_min, g2.y_max) == (0, 10, 1, 9)
    assert g2.seriler[0].kategori is not None                # X değerleri


def test_cubukta_eksen_basliklari_dosyada_yer_degistirir(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s.nesne_ekle(tam_grafik(s, "cubuk", x_baslik="Yatay", y_baslik="Dikey"))
    xlsx_io.kaydet(kitap, str(tmp_path / "a.xlsx"))
    c = load_workbook(tmp_path / "a.xlsx").worksheets[0]._charts[0]
    assert c.x_axis.title.tx.rich.p[0].r[0].t == "Dikey" and c.y_axis.title.tx.rich.p[0].r[0].t == "Yatay"
    g2 = xlsx_io.ac(str(tmp_path / "a.xlsx"))[0].sayfalar[0].nesneler[0]
    assert (g2.x_baslik, g2.y_baslik) == ("Yatay", "Dikey")


def test_pasta_her_dilime_acik_renk(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s.nesne_ekle(tam_grafik(s, "pasta"))
    xlsx_io.kaydet(kitap, str(tmp_path / "a.xlsx"))
    c = load_workbook(tmp_path / "a.xlsx").worksheets[0]._charts[0]
    assert len(c.series[0].dPt) == 3
    assert [getattr(d.spPr.solidFill.srgbClr, 'val', d.spPr.solidFill.srgbClr) for d in c.series[0].dPt][:2] == ["4472C4", "ED7D31"]


def test_ikinci_sayfa_ve_ozel_adli_sayfa_basvurusu(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s2 = kitap.sayfa_ekle("Satış 'Ö' Özeti")
    s2.ata(1, 1, Hucre(7.0))
    s2.ata(2, 1, Hucre(9.0))
    s.nesne_ekle(Grafik(seriler=(Seri(deger=Referans(s2.kimlik, Aralik(1, 1, 2, 1))),), konum=Konum(2, 2)))
    k2, _ = gidis_donus(kitap, tmp_path / "a.xlsx")
    g = k2.sayfalar[0].nesneler[0]
    assert g.seriler[0].deger == Referans(k2.sayfalar[1].kimlik, Aralik(1, 1, 2, 1))
    assert k2.sayfalar[1].ad == "Satış 'Ö' Özeti"


def test_silinmis_sayfaya_bakan_seri_yazilmaz(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s2 = kitap.sayfa_ekle("Gecici")
    s2.ata(1, 1, Hucre(1.0))
    g = Grafik(seriler=(Seri(deger=ref(s, 2, 2, 4, 2)), Seri(deger=Referans(s2.kimlik, Aralik(1, 1, 1, 1)))),
               konum=Konum(2, 2))
    s.nesne_ekle(g)
    kitap.sayfa_sil(s2)
    k2, rapor = gidis_donus(kitap, tmp_path / "a.xlsx")
    assert len(k2.sayfalar[0].nesneler[0].seriler) == 1 and not rapor.veri_kaybi


def test_seri_yoksa_grafik_yazilmaz(kitap, tmp_path):
    kitap.sayfalar[0].nesne_ekle(Grafik(konum=Konum(2, 2)))
    k2, rapor = gidis_donus(kitap, tmp_path / "a.xlsx")
    assert k2.sayfalar[0].nesneler == []
    with zipfile.ZipFile(tmp_path / "a.xlsx") as z:
        assert not any("charts" in n for n in z.namelist())


def test_resim_gidis_donus_baytlar_kayipsiz(kitap, tmp_path):
    s = kitap.sayfalar[0]
    veri = png((90, 40))
    s.nesne_ekle(Resim(veri=veri, ad="a.png", konum=Konum(4, 3, 5, 6), boyut=(180, 80)))
    k2, rapor = gidis_donus(kitap, tmp_path / "a.xlsx")
    r = k2.sayfalar[0].nesneler[0]
    assert isinstance(r, Resim) and r.veri == veri and r.konum == Konum(4, 3, 5, 6) and r.boyut == (180, 80)
    assert rapor.uyarilar == [] and not rapor.veri_kaybi


def test_grafik_ve_resim_birlikte_ve_kararli(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s.nesne_ekle(tam_grafik(s, "sutun", baslik="T"))
    s.nesne_ekle(Resim(veri=png(), konum=Konum(20, 2), boyut=(100, 50)))
    k2, _ = gidis_donus(kitap, tmp_path / "a.xlsx")
    k3, r3 = gidis_donus(k2, tmp_path / "b.xlsx")                 # kaydet → aç → kaydet → aç kararlı
    assert [type(n).__name__ for n in k3.sayfalar[0].nesneler] == ["Grafik", "Resim"]
    for a, b in zip(k2.sayfalar[0].nesneler, k3.sayfalar[0].nesneler):
        assert esit(a, b)
    assert not r3.veri_kaybi and r3.uyarilar == []


def test_tabcel_imi_yazilir(kitap, tmp_path):
    xlsx_io.kaydet(kitap, str(tmp_path / "a.xlsx"))
    with zipfile.ZipFile(tmp_path / "a.xlsx") as z:
        assert b"TabcelSurum" in z.read("docProps/custom.xml")


# ------------------------------------------------------------------ Excel / başka program dosyaları
def excel_dosyasi(yol, grafik3b=False, radar=False, resim=True, iki_hucre=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Veri"
    ws.append(["Ay", "Adet"])
    for i, ay in enumerate(["Ocak", "Şubat", "Mart"], 1):
        ws.append([ay, i * 10])
    g = BarChart()
    g.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    if iki_hucre:
        g.anchor = TwoCellAnchor(_from=AnchorMarker(col=3, row=1), to=AnchorMarker(col=8, row=11))
        ws.add_chart(g)
    else:
        ws.add_chart(g, "D2")
    if grafik3b:
        g3 = BarChart3D()
        g3.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        ws.add_chart(g3, "D20")
    if radar:
        r = RadarChart()
        r.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
        ws.add_chart(r, "D40")
    if resim:
        ws.add_image(Image(io.BytesIO(png())), "L2")
    wb.save(yol)
    return str(yol)


def test_excel_dosyasi_yuklenir_im_yok_uyarisi(tmp_path):
    k, rapor = xlsx_io.ac(excel_dosyasi(tmp_path / "e.xlsx"))
    assert [type(n).__name__ for n in k.sayfalar[0].nesneler] == ["Grafik", "Resim"]
    assert rapor.veri_kaybi and len(rapor.uyarilar) == 1
    assert "grafikler ve resimler" in rapor.uyarilar[0] and "uyarlandı" in rapor.uyarilar[0]


def test_yalniz_grafik_varsa_uyari_metni(tmp_path):
    _, rapor = xlsx_io.ac(excel_dosyasi(tmp_path / "e.xlsx", resim=False))
    assert "grafikler " in rapor.uyarilar[0] and "resim" not in rapor.uyarilar[0]


def test_uyarlanan_dosya_kaydedilince_im_gelir_uyari_kalkar(tmp_path):
    k, _ = xlsx_io.ac(excel_dosyasi(tmp_path / "e.xlsx"))
    k2, rapor2 = gidis_donus(k, tmp_path / "t.xlsx")
    assert rapor2.uyarilar == [] and not rapor2.veri_kaybi and len(k2.sayfalar[0].nesneler) == 2


def test_desteklenmeyen_turler_atlanir_ve_uyarilir(tmp_path):
    k, rapor = xlsx_io.ac(excel_dosyasi(tmp_path / "e.xlsx", grafik3b=True, radar=True, resim=False))
    assert len(k.sayfalar[0].nesneler) == 1                                 # yalnız sütun grafiği
    metin = " ".join(rapor.uyarilar)
    assert "2 grafik desteklenmiyor" in metin and "kaydederseniz" in metin and rapor.veri_kaybi


def test_iki_hucre_capasinda_boyut_hucre_olculerinden(tmp_path):
    k, _ = xlsx_io.ac(excel_dosyasi(tmp_path / "e.xlsx", iki_hucre=True, resim=False))
    g = k.sayfalar[0].nesneler[0]
    s = k.sayfalar[0]
    assert g.konum == Konum(2, 4)
    assert g.boyut == (sum(s.genislik(c) for c in range(4, 9)), sum(s.yukseklik(r) for r in range(2, 12)))


def test_cok_aralikli_seri_atlanir_uyarilir(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "V"
    for i in range(1, 6):
        ws.cell(row=i, column=1, value=i)
    g = BarChart()
    g.series.append(Series(Reference(ws, min_col=1, min_row=1, max_row=3), title="a"))
    g.series[0].val.numRef.f = "(V!$A$1:$A$2,V!$A$4:$A$5)"                   # çok aralıklı
    g.series.append(Series(Reference(ws, min_col=1, min_row=1, max_row=5), title="b"))
    ws.add_chart(g, "C2")
    wb.save(tmp_path / "c.xlsx")
    k, rapor = xlsx_io.ac(str(tmp_path / "c.xlsx"))
    assert len(k.sayfalar[0].nesneler[0].seriler) == 1
    assert "1 grafik serisi okunamadı" in " ".join(rapor.uyarilar) and rapor.veri_kaybi


def test_im_var_ama_okunamayan_grafik_yine_uyarir(kitap, tmp_path):
    s = kitap.sayfalar[0]
    s.nesne_ekle(tam_grafik(s))
    xlsx_io.kaydet(kitap, str(tmp_path / "a.xlsx"))
    with zipfile.ZipFile(tmp_path / "a.xlsx") as z, zipfile.ZipFile(tmp_path / "c.xlsx", "w") as o:
        for ad in z.namelist():
            o.writestr(ad, z.read(ad))
        o.writestr("xl/charts/chart2.xml", "<c/>")                             # dışarıdan eklenen bozuk grafik
    _, rapor = xlsx_io.ac(str(tmp_path / "c.xlsx"))
    assert rapor.veri_kaybi and "1 grafik desteklenmiyor" in " ".join(rapor.uyarilar)


def test_kullanilmayan_orfan_medya_uyari_vermez(kitap, tmp_path):
    xlsx_io.kaydet(kitap, str(tmp_path / "a.xlsx"))
    with zipfile.ZipFile(tmp_path / "a.xlsx") as z, zipfile.ZipFile(tmp_path / "b.xlsx", "w") as o:
        for ad in z.namelist():
            o.writestr(ad, z.read(ad))
        o.writestr("xl/media/image9.png", png())                              # hiçbir çizimin başvurmadığı resim
    _, rapor = xlsx_io.ac(str(tmp_path / "b.xlsx"))
    assert not rapor.veri_kaybi


def test_sekil_iceren_cizim_hala_kaybolur_uyarisi_verir(kitap, tmp_path):
    xlsx_io.kaydet(kitap, str(tmp_path / "a.xlsx"))
    cizim = ('<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
             '<xdr:twoCellAnchor><xdr:sp></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>')
    with zipfile.ZipFile(tmp_path / "a.xlsx") as z, zipfile.ZipFile(tmp_path / "b.xlsx", "w") as o:
        for ad in z.namelist():
            o.writestr(ad, z.read(ad))
        o.writestr("xl/drawings/drawing1.xml", cizim)
    rapor = xlsx_io.tara(str(tmp_path / "b.xlsx"))
    assert "şekil" in " ".join(rapor.uyarilar) and rapor.veri_kaybi


def test_ornek_grafikler_xlsx_uyarisiz_ve_kararli(tmp_path):
    """6j: `testler/ornekler/grafikler.xlsx` (Tabcel'in yazdığı, 13 grafik + resim) uyarısız açılır, gidiş-dönüşte aynı kalır."""
    import os
    yol = os.path.join(os.path.dirname(__file__), "ornekler", "grafikler.xlsx")
    k, rapor = xlsx_io.ac(yol)
    n = k.sayfalar[0].nesneler
    assert len(n) == 14 and not rapor.veri_kaybi and not rapor.uyarilar
    xlsx_io.kaydet(k, str(tmp_path / "b.xlsx"))
    k2, rapor2 = xlsx_io.ac(str(tmp_path / "b.xlsx"))
    assert [x.degistir(kimlik=0) for x in k2.sayfalar[0].nesneler] == [x.degistir(kimlik=0) for x in n]
    assert not rapor2.veri_kaybi
