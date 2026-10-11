"""Eski Excel (.xls, BIFF) okuma — Faz 8a. Yalnız OKUMA (kaydetmek .xlsx'e yapılır); xlrd (BSD) kullanır.

Okunanlar: sayfalar (gizli sayfa görünür yapılır + uyarı), hücre değerleri (sayı, metin, mantıksal, hata, tarih →
seri sayı), hücre biçimi (yazı tipi, dolgu, kenarlık, hizalama, sayı biçimi), birleşimler, sütun genişliği/satır
yüksekliği/gizleme, dondurma, kılavuz çizgisi, notlar. **Okunamayanlar** (uyarı verilir): formüller (xlrd formül
metnini vermez; hücreler son hesaplanan DEĞERLERLE açılır), grafik/resim/makro/pivot, koşullu biçim, veri doğrulama,
tanımlı adlar, filtre. `.xls` adıyla kaydedilmiş gerçek bir .xlsx dosyası doğrudan `xlsx_io.ac`a yönlenir."""
import os
from typing import Dict, Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color

from ..model import olcu
from ..model.adres import MAX_COL, MAX_ROW, Aralik
from ..model.deger import Hata, hata_bul
from ..model.notlar import Not
from ..model.sayfa import CalismaKitabi, Hucre
from . import xlsx_io
from .xlsx_io import Rapor, XlsxHatasi, _stil_oku

_YATAY = {0: None, 1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous"}
_DIKEY = {0: "top", 1: "center", 2: "bottom", 3: "justify"}
_CIZGI = {1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double", 7: "hair",
          8: "mediumDashed", 9: "dashDot", 10: "mediumDashDot", 11: "dashDotDot", 12: "mediumDashDotDot",
          13: "slantDashDot"}
_ALTI = {1: "single", 2: "double", 33: "singleAccounting", 34: "doubleAccounting"}
_OLE = b"\xd0\xcf\x11\xe0"


def _renk(kitap, indeks) -> Optional[Color]:
    """xlrd renk indeksi → openpyxl Color; sistem/otomatik renkler (0x40, 0x41, 0x7FFF) None."""
    rgb = kitap.colour_map.get(indeks) if isinstance(indeks, int) else None
    return Color(rgb="FF%02X%02X%02X" % rgb) if rgb else None


def _yazi_rengi(kitap, indeks) -> Optional[Color]:
    """Siyah yazı = varsayılan (her hücreye açık renk yazılmasın)."""
    c = _renk(kitap, indeks)
    return None if c is not None and c.rgb == "FF000000" else c


def _stil_nesneleri(kitap, xf, temel_font):
    """xlrd XF kaydı → (Font, PatternFill, Border, Alignment, sayı biçimi) — `xlsx_io._stil_oku`nun girdisi."""
    f = kitap.font_list[xf.font_index]
    ad = f.name if f.name and f.name != temel_font.name else "Calibri"       # çalışma kitabının varsayılan yazı tipi → bizim varsayılan
    boyut = f.height / 20 if f.height != temel_font.height else 11.0
    font = Font(name=ad, sz=boyut, b=bool(f.bold), i=bool(f.italic), strike=bool(f.struck_out),
                u=_ALTI.get(f.underline_type), color=_yazi_rengi(kitap, f.colour_index))
    dolgu = PatternFill()
    if xf.background.fill_pattern == 1:
        c = _renk(kitap, xf.background.pattern_colour_index)
        if c is not None:
            dolgu = PatternFill("solid", fgColor=c, bgColor=c)

    def kenar(cizgi, renk):
        return Side(style=_CIZGI.get(cizgi), color=_renk(kitap, renk) or Color(rgb="FF000000")) if cizgi else Side()
    b = xf.border
    kenarlik = Border(top=kenar(b.top_line_style, b.top_colour_index), bottom=kenar(b.bottom_line_style, b.bottom_colour_index),
                      left=kenar(b.left_line_style, b.left_colour_index), right=kenar(b.right_line_style, b.right_colour_index))
    a = xf.alignment
    hiza = Alignment(horizontal=_YATAY.get(a.hor_align), vertical=_DIKEY.get(a.vert_align, "bottom"),
                     indent=a.indent_level, wrap_text=bool(a.text_wrapped))
    bicim = kitap.format_map[xf.format_key].format_str if xf.format_key in kitap.format_map else "General"
    return font, dolgu, kenarlik, hiza, bicim


def ac(yol: str) -> "tuple[CalismaKitabi, Rapor]":
    """`yol`daki .xls dosyasını yeni bir kitaba okur (bkz. modül açıklaması)."""
    try:
        with open(yol, "rb") as h:
            bas = h.read(4)
    except OSError as e:
        raise XlsxHatasi(f"Dosya açılamadı: {e.strerror or e}") from e
    if bas[:2] == b"PK":                                       # uzantısı .xls ama aslında .xlsx (openpyxl uzantıya bakar)
        import shutil
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            kopya = os.path.join(d, "kitap.xlsx")
            shutil.copyfile(yol, kopya)
            return xlsx_io.ac(kopya)
    if bas != _OLE:
        raise XlsxHatasi("Dosya geçerli bir .xls değil (bozuk ya da farklı biçim).")
    import xlrd
    try:
        kt = xlrd.open_workbook(yol, formatting_info=True, on_demand=False)
    except xlrd.XLRDError as e:
        mesaj = str(e).lower()
        if "encrypt" in mesaj or "password" in mesaj:
            raise XlsxHatasi("Bu .xls dosyası parolayla korunuyor; şimdilik açılamıyor.") from e
        raise XlsxHatasi(f"Dosya okunamadı: {e}") from e
    except (OSError, ValueError, IndexError, AssertionError, KeyError) as e:     # bozuk BIFF kayıtları çeşitli hata verir
        raise XlsxHatasi(f"Dosya okunamadı (bozuk olabilir): {e}") from e

    rapor = Rapor()
    kitap = CalismaKitabi()
    havuz = kitap.stil_havuzu
    temel_font = kt.font_list[kt.xf_list[0].font_index] if kt.xf_list else None
    stil_onb: Dict[int, int] = {}
    kayma = 1462 if kt.datemode == 1 else 0                    # 1904 sistemi: seri sayıyı 1900 sistemine çevir
    if kayma:
        rapor.ekle("Bu dosya 1904 tarih sistemini kullanıyor; tarihler 1900 sistemine çevrildi.")
    gizli = 0
    for i in range(kt.nsheets):
        ws = kt.sheet_by_index(i)
        s = kitap.sayfalar[0] if i == 0 else kitap.sayfa_ekle(ws.name)
        if i == 0:
            kitap.yeniden_adlandir(s, ws.name)
        if ws.visibility != 0:
            gizli += 1
        yeni, stiller, notlar = {}, {}, {}
        for r in range(min(ws.nrows, MAX_ROW)):
            for c in range(min(ws.ncols, MAX_COL)):
                tur = ws.cell_type(r, c)
                if tur in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
                    yeni[(r + 1, c + 1)] = Hucre(float(ws.cell_value(r, c)) + (kayma if tur == xlrd.XL_CELL_DATE else 0))
                elif tur == xlrd.XL_CELL_TEXT:
                    v = ws.cell_value(r, c)
                    if v != "":
                        yeni[(r + 1, c + 1)] = Hucre(v)
                elif tur == xlrd.XL_CELL_BOOLEAN:
                    yeni[(r + 1, c + 1)] = Hucre(bool(ws.cell_value(r, c)))
                elif tur == xlrd.XL_CELL_ERROR:
                    metin = xlrd.error_text_from_code.get(ws.cell_value(r, c), "#VALUE!")
                    yeni[(r + 1, c + 1)] = Hucre(hata_bul(metin) or Hata(metin))
                xi = ws.cell_xf_index(r, c)
                k = stil_onb.get(xi)
                if k is None:
                    xf = kt.xf_list[xi]
                    k = stil_onb[xi] = havuz.kimlik(_stil_oku(*_stil_nesneleri(kt, xf, temel_font))) if temel_font else 0
                if k:
                    stiller[(r + 1, c + 1)] = k
        s.toplu_yaz(yeni)
        for (r, c), n in ws.cell_note_map.items():
            metin = (n.text or "").replace("\r\n", "\n").replace("\r", "\n")
            if metin:
                notlar[(r + 1, c + 1)] = Not(metin, (n.author or "").strip())
        if notlar:
            s.notlari_yaz(notlar)
        for r, ri in ws.rowinfo_map.items():
            if ri.height_mismatch and ri.height:
                s.satir_yukseklik[r + 1] = olcu.yukseklik_sinirla(olcu.punto_px(ri.height / 20))
            if ri.hidden:
                s.gizli_satir.add(r + 1)
        for c, ci in ws.colinfo_map.items():
            if c >= MAX_COL:
                continue
            if ci.width:
                s.sutun_genislik[c + 1] = olcu.genislik_sinirla(olcu.karakter_px(ci.width / 256))
            if ci.hidden:
                s.gizli_sutun.add(c + 1)
        s.bicim_yaz(stiller, {}, {})
        birlesimler = []
        for rlo, rhi, clo, chi in ws.merged_cells:
            a = Aralik(rlo + 1, clo + 1, rhi, chi)
            if (a.ilk_satir, a.ilk_sutun) == (a.son_satir, a.son_sutun):
                continue
            if any(not (a.son_satir < x.ilk_satir or a.ilk_satir > x.son_satir or a.son_sutun < x.ilk_sutun
                        or a.ilk_sutun > x.son_sutun) for x in birlesimler):
                continue
            birlesimler.append(a)
        s.birlesim_yukle(birlesimler)
        dondur = None
        if ws.panes_are_frozen and (ws.horz_split_pos or ws.vert_split_pos):
            dondur = (int(ws.horz_split_pos) + 1, int(ws.vert_split_pos) + 1)
            if dondur == (1, 1):
                dondur = None
        s.gorunum_yaz(dondur=dondur, kilavuz=bool(ws.show_grid_lines))
    if gizli:
        rapor.ekle(f"{gizli} gizli sayfa görünür yapıldı.")
    rapor.ekle("Bu bir eski .xls dosyası: formüller okunamaz, hücreler son hesaplanan değerleriyle açıldı. "
               "Grafik, resim, makro, koşullu biçim, veri doğrulama ve tanımlı adlar da okunmaz. "
               "Kaydederken .xlsx olarak yeni dosya oluşturulur.")
    return kitap, rapor
